#!/usr/bin/env python3
"""Local PDF tools with no-clobber publication and explicit document preservation.

Coordinates are PDF points, measured from the upper-left corner of an unrotated
page. Page selectors are one-based. Every operation writes a new destination.
"""
from __future__ import annotations
import argparse
import base64
import ctypes
import difflib
import html
import importlib.util
import io
import json
import math
import os
from pathlib import Path
import re
import secrets
import signal
import shutil
import subprocess
import sys
import tempfile
import fitz

TOOLS = {
    'merge':'Juntar PDF', 'split':'Dividir PDF', 'compress':'Comprimir PDF',
    'word':'PDF para Word', 'powerpoint':'PDF para PowerPoint', 'excel':'PDF para Excel',
    'office':'Office para PDF', 'text':'Adicionar texto', 'edit':'Editar texto',
    'image':'Inserir imagem', 'shape':'Desenhar forma', 'jpg':'PDF para JPG', 'png':'PDF para PNG',
    'odt':'ODT para PDF', 'txt':'TXT para PDF', 'rtf':'RTF para PDF', 'epub':'EPUB para PDF', 'csv':'CSV para PDF',
    'flatten':'Achatar PDF', 'delete-pages':'Excluir páginas', 'extract-pages':'Extrair páginas',
    'images':'Imagens para PDF', 'signature':'Assinatura visual', 'watermark':'Marca d’água',
    'rotate':'Rodar PDF', 'html':'HTML para PDF', 'unlock':'Desbloquear PDF',
    'protect':'Proteger PDF', 'organize':'Organizar PDF', 'pdfa':'PDF para PDF/A',
    'repair':'Reparar PDF', 'numbers':'Números de página', 'ocr':'OCR PDF',
    'compare':'Comparar texto', 'compare-visual':'Comparar visualmente',
    'redact':'Remover conteúdo', 'crop':'Recortar PDF', 'forms':'Inspecionar formulários',
    'form-create':'Criar formulários', 'form-fill':'Preencher formulários', 'form-detect':'Detectar campos',
    'layout-extract':'Extrair blocos de layout', 'translate-layout':'Aplicar tradução ao layout',
    'markdown':'PDF para Markdown', 'workflow':'Executar fluxo'
}
ALIASES = {'edit_text':'edit', 'form_create':'form-create', 'form_fill':'form-fill', 'compare_visual':'compare-visual','delete_pages':'delete-pages','extract_pages':'extract-pages','delete':'delete-pages','extract':'extract-pages'}
PDF_OPERATIONS = {'merge','compress','office','text','edit','image','shape','images','signature','watermark','rotate','html','unlock','protect','organize','pdfa','repair','numbers','ocr','redact','crop','form-create','form-fill','translate-layout','odt','rtf','txt','epub','csv','flatten','delete-pages','extract-pages'}
MULTI_INPUT = {'merge','images','compare','compare-visual','organize'}
DIRECTORY_OUTPUT = {'split','jpg','png'}


def open_pdf(path, password=''):
    doc = fitz.open(path)
    if doc.needs_pass and not doc.authenticate(password):
        doc.close()
        raise ValueError('Senha necessária ou incorreta.')
    return doc


def save(doc, output, **kwargs):
    _separate_page_aliases(doc)
    options = dict(garbage=4, deflate=True, encryption=fitz.PDF_ENCRYPT_KEEP)
    options.update(kwargs)
    doc.save(str(output), **options)


def number(value, name, minimum=None, maximum=None):
    if isinstance(value,bool): raise ValueError(f'{name}: informe um número, não um booleano.')
    try:
        result = float(value)
    except (TypeError, ValueError):
        raise ValueError(f'{name}: informe um número válido.') from None
    if not math.isfinite(result) or (minimum is not None and result < minimum) or (maximum is not None and result > maximum):
        raise ValueError(f'{name}: valor fora do intervalo permitido.')
    return result


def integer(value,name,minimum=None,maximum=None):
    result=number(value,name,minimum,maximum)
    if not result.is_integer(): raise ValueError(f'{name}: informe um número inteiro.')
    return int(result)


def _separate_page_aliases(doc):
    """Duplicated page-tree references must not share later mutations."""
    indices=[]; seen=set(); initial=len(doc)
    for index in range(initial):
        xref=doc[index].xref
        if xref not in seen:
            indices.append(index); seen.add(xref); continue
        with fitz.open() as copy:
            copy.insert_pdf(doc,from_page=index,to_page=index)
            doc.insert_pdf(copy)
        indices.append(len(doc)-1)
    if len(doc)!=initial: doc.select(indices)


def color(value, default=(0.2, 0.2, 0.25)):
    if value is None:
        return default
    if isinstance(value, str):
        value = value.removeprefix('#')
        if not re.fullmatch(r'[0-9a-fA-F]{6}', value):
            raise ValueError('Cor inválida; use #RRGGBB.')
        return tuple(int(value[i:i+2],16)/255 for i in (0,2,4))
    if not isinstance(value,(list,tuple)) or len(value) != 3:
        raise ValueError('Cor inválida.')
    return tuple(number(v,'Cor',0,1) for v in value)


def parse_pages(selector, count, allow_duplicates=False):
    if selector is None or str(selector).strip() in ('','all','*'):
        return list(range(count))
    result = []
    tokens = selector if isinstance(selector,list) else str(selector).split(',')
    for token in tokens:
        value = str(token).strip()
        match = re.fullmatch(r'(\d+)(?:\s*-\s*(\d+))?',value)
        if not match:
            raise ValueError('Páginas inválidas. Use, por exemplo, 1,3-5.')
        first = int(match[1]); last = int(match[2] or first)
        for page in range(first,last+(1 if last >= first else -1),1 if last >= first else -1):
            if page < 1 or page > count:
                raise ValueError(f'Página {page} fora do documento ({count} páginas).')
            if allow_duplicates or page-1 not in result:
                result.append(page-1)
    if not result:
        raise ValueError('Selecione ao menos uma página.')
    return result


def _area(rect):
    return max(0,rect.width)*max(0,rect.height)


def _boolean(value):
    if isinstance(value,bool): return value
    if value in (0,1): return bool(value)
    if isinstance(value,str):
        normalized=value.strip().lower()
        if normalized in ('true','yes','on','1'): return True
        if normalized in ('false','no','off','0',''): return False
    raise ValueError('Valor de caixa de seleção inválido; use true ou false.')


def _page_bounds(page):
    return page.rect * page.derotation_matrix


def rectangle(options, page, default=None):
    bounds = _page_bounds(page)
    if 'rect' in options:
        values = options['rect']
        if not isinstance(values,(list,tuple)) or len(values) != 4:
            raise ValueError('rect deve conter [x0,y0,x1,y1].')
        rect = fitz.Rect([number(v,'Coordenada') for v in values])
    elif any(key in options for key in ('x','y','width','height')):
        x = number(options.get('x',36),'x'); y = number(options.get('y',72),'y')
        rect = fitz.Rect(x,y,x+number(options.get('width',200),'Largura',1),y+number(options.get('height',80),'Altura',1))
    else:
        rect = fitz.Rect(default or [36,72,min(236,bounds.width-36),min(152,bounds.height-36)])
    if rect.is_empty or rect.is_infinite or not bounds.contains(rect):
        raise ValueError('A área deve estar inteiramente dentro da página e ter tamanho positivo.')
    return rect


def _encrypted_output(doc, password):
    # New documents cannot use ENCRYPT_KEEP because they have no source crypt dict.
    if doc.metadata.get('encryption'):
        if not password:
            raise ValueError('Informe a senha para preservar a proteção da cópia.')
        return dict(encryption=fitz.PDF_ENCRYPT_AES_256, user_pw=password,
                    owner_pw=secrets.token_urlsafe(24), permissions=doc.permissions)
    return {}


def _publish(payload, output):
    """Publish complete output without ever replacing a concurrently created path."""
    if payload.is_file():
        os.link(payload, output)  # atomic EEXIST on destination, including symlinks
        payload.unlink()
        return
    libc = ctypes.CDLL(None, use_errno=True)
    if sys.platform == 'darwin':
        rename = libc.renamex_np
        rename.argtypes = [ctypes.c_char_p,ctypes.c_char_p,ctypes.c_uint]
        status = rename(os.fsencode(payload),os.fsencode(output),4)  # RENAME_EXCL
    elif sys.platform.startswith('linux') and hasattr(libc,'renameat2'):
        rename = libc.renameat2
        rename.argtypes = [ctypes.c_int,ctypes.c_char_p,ctypes.c_int,ctypes.c_char_p,ctypes.c_uint]
        status = rename(-100,os.fsencode(payload),-100,os.fsencode(output),1)  # RENAME_NOREPLACE
    else:
        raise RuntimeError('Publicação atômica de diretórios exige macOS ou Linux compatível.')
    if status:
        error = ctypes.get_errno()
        raise OSError(error,os.strerror(error),str(output))


def _external(command, timeout=180):
    """Own a separate process group so cancel never leaves converters running."""
    try:
        process = subprocess.Popen([str(v) for v in command],stdout=subprocess.PIPE,stderr=subprocess.PIPE,stdin=subprocess.DEVNULL,text=True,start_new_session=True)
    except FileNotFoundError:
        raise RuntimeError(f'Dependência indisponível: {command[0]}. Execute o instalador completo.') from None
    try:
        stdout,stderr = process.communicate(timeout=timeout)
    except BaseException:
        try: os.killpg(process.pid,signal.SIGTERM)
        except ProcessLookupError: pass
        try:
            process.communicate(timeout=3)
        except subprocess.TimeoutExpired:
            pass
        finally:
            # The leader can exit while a descendant ignores TERM and closes its
            # pipes. Always finish the owned group, even after communicate returns.
            try: os.killpg(process.pid,signal.SIGKILL)
            except ProcessLookupError: pass
            process.communicate()
        raise
    result = subprocess.CompletedProcess(command,process.returncode,stdout,stderr)
    if result.returncode:
        message = (result.stderr or result.stdout).strip()[-2500:]
        raise RuntimeError(f'{Path(command[0]).name} falhou ({result.returncode}): {message}')
    return result


def run(operation, inputs, output, options=None):
    options = {} if options is None else options
    if not isinstance(options,dict):
        raise ValueError('As opções devem ser um objeto JSON.')
    if not isinstance(operation,str): raise ValueError('Informe uma operação válida.')
    operation = ALIASES.get(operation,operation)
    if operation == 'compare' and options.get('mode') == 'visual': operation = 'compare-visual'
    if operation == 'jpg' and options.get('format') == 'png': operation = 'png'
    if operation not in TOOLS: raise ValueError('Operação desconhecida: '+str(operation))
    if inputs is None or isinstance(inputs,(str,Path)): raise ValueError('As entradas devem ser uma lista de arquivos.')
    inputs = [str(path) for path in inputs]
    if not inputs: raise ValueError('Selecione ao menos um arquivo.')
    if operation not in MULTI_INPUT | {'workflow'} and len(inputs) != 1:
        raise ValueError('Esta operação aceita um arquivo por vez.')
    if operation in ('compare','compare-visual') and len(inputs) != 2:
        raise ValueError('Selecione dois PDFs para comparar.')
    requested = Path(output).expanduser().absolute()
    output = requested.parent.resolve() / requested.name
    if os.path.lexists(output) or any(output.resolve() == Path(p).expanduser().resolve() for p in inputs if not p.startswith(('https://','http://'))):
        raise ValueError('O destino já existe. Escolha um novo nome para preservar os arquivos.')
    output.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.pdf-output-',dir=output.parent) as temporary:
        payload = Path(temporary)/output.name
        _execute(operation,inputs,payload,options)
        if not payload.exists(): raise RuntimeError('A operação não produziu o arquivo esperado.')
        if payload.is_file():
            with payload.open('rb') as handle: os.fsync(handle.fileno())
        try: _publish(payload,output)
        except FileExistsError:
            raise ValueError('Outro processo criou o destino. Nenhum arquivo foi substituído.') from None
    return str(output)


def _execute(operation, inputs, output, options):
    password = options.get('password','')
    if operation=='organize' and len(inputs)==1 and 'input_passwords' in options:
        credentials=options['input_passwords']
        if not isinstance(credentials,list) or len(credentials)!=1 or not isinstance(credentials[0],str): raise ValueError('input_passwords deve conter uma senha por arquivo de entrada.')
        password=credentials[0]
    if not isinstance(password,str): raise ValueError('A senha deve ser uma string.')
    if operation == 'workflow':
        steps = options.get('steps',[])
        if not isinstance(steps,list) or not 1 <= len(steps) <= 64:
            raise ValueError('O fluxo precisa de 1 a 64 etapas.')
        for i, step in enumerate(steps):
            if not isinstance(step,dict) or step.get('operation') not in TOOLS or step['operation'] == 'workflow':
                raise ValueError(f'Etapa {i+1} inválida; fluxos recursivos não são permitidos.')
            if not isinstance(step.get('options',{}),dict): raise ValueError(f'As opções da etapa {i+1} devem ser um objeto JSON.')
            if i < len(steps)-1 and step['operation'] not in PDF_OPERATIONS:
                raise ValueError('Etapas intermediárias precisam produzir um PDF único.')
            if i > 0 and step['operation'] in ('merge','images','compare','compare-visual'):
                raise ValueError('Esta operação só pode ser a primeira etapa do fluxo.')
        with tempfile.TemporaryDirectory() as tmp:
            current = inputs
            for i,step in enumerate(steps):
                target = output if i == len(steps)-1 else Path(tmp)/f'{i}.pdf'
                step_options = dict(step.get('options',{}))
                if password and 'password' not in step_options: step_options['password'] = password
                run(step['operation'],current,target,step_options)
                password = step_options.get('output_password',step_options.get('password',password))
                if step['operation'] in ('merge','organize') and not password:
                    password=next((value for value in step_options.get('input_passwords',[]) if value),'')
                if step['operation'] == 'protect': password = step_options.get('new_password','')
                if step['operation'] == 'unlock': password = ''
                current = [str(target)]
        return
    if operation == 'txt':
        _text_file_pdf(inputs[0],output,options); return
    if operation == 'epub':
        width,height,margin = _paper_geometry(options)
        fontsize = number(options.get('font_size',11),'Tamanho da fonte',4,72)
        with fitz.open(inputs[0],filetype='epub') as book:
            if not book.is_reflowable: raise ValueError('Selecione um EPUB válido e sem DRM.')
            book.layout(width=width-2*margin,height=height-2*margin,fontsize=fontsize)
            with fitz.open('pdf',book.convert_to_pdf()) as converted, fitz.open() as result:
                for i in range(len(converted)):
                    page=result.new_page(width=width,height=height)
                    page.show_pdf_page(page.rect+(margin,margin,-margin,-margin),converted,i)
                if not len(result): raise ValueError('O EPUB não contém páginas convertíveis.')
                save(result,output)
        return
    if operation in ('office','odt','rtf','csv'):
        executable = shutil.which('soffice') or '/Applications/LibreOffice.app/Contents/MacOS/soffice'
        with tempfile.TemporaryDirectory() as tmp:
            source = inputs[0]
            if operation == 'csv' or Path(source).suffix.lower() == '.csv':
                source = str(_csv_workbook(source,Path(tmp)/'source.xlsx',options))
                options = dict(options); options.setdefault('fit_width',True)
            if Path(source).suffix.lower() == '.xlsx' and (options.get('fit_width') or options.get('landscape') or options.get('orientation')):
                from openpyxl import load_workbook
                book = load_workbook(source)
                for sheet in book:
                    sheet.page_setup.orientation = 'landscape' if options.get('landscape') or options.get('orientation') == 'landscape' else 'portrait'
                    if options.get('fit_width'):
                        sheet.sheet_properties.pageSetUpPr.fitToPage = True
                        sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
                source = str(Path(tmp)/'input.xlsx'); book.save(source); book.close()
            _external([executable,'-env:UserInstallation='+(Path(tmp)/'profile').as_uri(),'--headless','--convert-to','pdf','--outdir',tmp,source])
            generated = Path(tmp)/(Path(source).stem+'.pdf')
            if not generated.exists(): raise RuntimeError('LibreOffice não produziu um PDF.')
            with fitz.open(generated) as doc:
                if len(doc) == 0: raise RuntimeError('LibreOffice produziu um documento vazio.')
            shutil.copyfile(generated,output)
        return
    if operation in ('ocr','pdfa'):
        with open_pdf(inputs[0],password) as doc, tempfile.TemporaryDirectory() as tmp:
            protection = _encrypted_output(doc,password)
            if protection and operation == 'pdfa':
                raise ValueError('PDF/A não permite criptografia. Desbloqueie uma cópia explicitamente antes de converter.')
            source = inputs[0]
            if protection:
                source = str(Path(tmp)/'source.pdf'); save(doc,source,encryption=fitz.PDF_ENCRYPT_NONE)
            generated = Path(tmp)/'result.pdf'
            _external(['ocrmypdf','--output-type','pdfa' if operation == 'pdfa' else 'pdf','--skip-text','--language',options.get('language','por+eng'),source,generated],timeout=1800)
            with open_pdf(generated) as result:
                if protection: save(result,output,**protection)
                else: shutil.copyfile(generated,output)
        return
    if operation == 'html':
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            browser = p.chromium.launch(executable_path=os.getenv('PDF_CHROMIUM_EXECUTABLE') or None)
            try:
                page = browser.new_page()
                if inputs[0].startswith(('https://','http://')):
                    page.goto(inputs[0],wait_until='networkidle',timeout=60000)
                else:
                    page.set_content(Path(inputs[0]).read_text(encoding='utf-8'),wait_until='networkidle',timeout=60000)
                page.pdf(path=str(output),format=options.get('paper','A4'),print_background=True,landscape=bool(options.get('landscape',False)))
            finally: browser.close()
        return
    if operation in ('merge','images') or (operation=='organize' and len(inputs)>1):
        protection = {}
        input_passwords=options.get('input_passwords',[password]*len(inputs))
        if not isinstance(input_passwords,list) or len(input_passwords)!=len(inputs) or any(not isinstance(value,str) for value in input_passwords):
            raise ValueError('input_passwords deve conter uma senha por arquivo de entrada, na mesma ordem.')
        output_password=options.get('output_password'); protected_passwords=set()
        if output_password is not None and (not isinstance(output_password,str) or not 8<=len(output_password)<=40):
            raise ValueError('A senha do resultado deve conter de 8 a 40 caracteres.')
        with fitz.open() as result:
            for index,path in enumerate(inputs):
                source_password=input_passwords[index]
                with open_pdf(path,source_password) as source:
                    source_protection = _encrypted_output(source,source_password) if source.is_pdf else {}
                    if source_protection:
                        protected_passwords.add(source_password)
                        if len(protected_passwords)>1 and output_password is None:
                            raise ValueError('Os PDFs usam senhas diferentes. Defina output_password para proteger o resultado.')
                        if output_password is not None: source_protection['user_pw']=output_password
                    if source_protection:
                        if protection: source_protection['permissions'] &= protection['permissions']
                        protection = source_protection
                    if operation == 'images':
                        with fitz.open('pdf',source.convert_to_pdf()) as converted:
                            for i in range(len(converted)):
                                if any(k in options for k in ('paper','margin','landscape','orientation')):
                                    width,height = fitz.paper_size(options.get('paper','a4'))
                                    if width <= 0 or height <= 0: raise ValueError('Papel desconhecido.')
                                    if options.get('landscape') or options.get('orientation') == 'landscape': width,height = height,width
                                    page = result.new_page(width=width,height=height)
                                    margin = number(options.get('margin',20),'Margem',0,min(width,height)/3)
                                    page.show_pdf_page(page.rect+(margin,margin,-margin,-margin),converted,i)
                                else: result.insert_pdf(converted,from_page=i,to_page=i)
                    else: result.insert_pdf(source)
            if not len(result): raise ValueError('Nenhuma página encontrada.')
            if operation=='organize':
                if not options.get('pages'): raise ValueError('Informe a nova ordem das páginas concatenadas.')
                result.select(parse_pages(options['pages'],len(result),allow_duplicates=True))
            if output_password and not protection:
                protection=dict(encryption=fitz.PDF_ENCRYPT_AES_256,user_pw=output_password,owner_pw=secrets.token_urlsafe(24))
            save(result,output,**protection)
        return
    with open_pdf(inputs[0],password) as doc:
        if not doc.is_pdf or not len(doc): raise ValueError('Selecione um PDF com ao menos uma página.')
        if operation in PDF_OPERATIONS: _separate_page_aliases(doc)
        pages = parse_pages(options.get('pages'),len(doc),allow_duplicates=operation in ('organize','extract-pages'))
        if operation == 'split':
            output.mkdir()
            groups = options.get('groups')
            if groups is not None and (not isinstance(groups,list) or not groups or any(not str(group).strip() for group in groups)):
                raise ValueError('Os grupos devem ser uma lista não vazia de intervalos de páginas.')
            selected_groups = [parse_pages(group,len(doc),allow_duplicates=True) for group in groups] if groups else [[p] for p in pages]
            for i, group in enumerate(selected_groups):
                with fitz.open() as part:
                    for p in group: part.insert_pdf(doc,from_page=p,to_page=p)
                    save(part,output/f'parte-{i+1:04}.pdf',**_encrypted_output(doc,password))
        elif operation in ('jpg','png'):
            _raster_export(doc,pages,output,operation,options)
        elif operation == 'layout-extract':
            output.write_text(json.dumps(_layout(doc,pages),ensure_ascii=False,indent=2),encoding='utf-8')
        elif operation == 'translate-layout':
            _translate_layout(doc,options); save(doc,output)
        elif operation == 'form-detect':
            output.write_text(json.dumps(dict(heuristic=True,review_required=True,fields=_detect_forms(doc,pages)),ensure_ascii=False,indent=2),encoding='utf-8')
        elif operation == 'forms':
            content = _form_records(doc,pages)
            output.write_text(json.dumps(content,ensure_ascii=False,indent=2),encoding='utf-8')
        elif operation == 'markdown':
            _markdown(doc,pages,output)
        elif operation in ('compare','compare-visual'):
            with open_pdf(inputs[1],options.get('second_password',password)) as other:
                if operation == 'compare-visual': _visual_compare(doc,other,output)
                else:
                    first = ''.join(p.get_text() for p in doc).splitlines(True)
                    second = ''.join(p.get_text() for p in other).splitlines(True)
                    diff = ''.join(difflib.unified_diff(first,second,fromfile=Path(inputs[0]).name,tofile=Path(inputs[1]).name))
                    output.write_text(diff or 'Nenhuma diferença textual. Imagens e layout não foram comparados.\n',encoding='utf-8')
        elif operation == 'word': _word(doc,pages,output)
        elif operation == 'powerpoint': _powerpoint(doc,pages,output,options)
        elif operation == 'excel': _excel(doc,pages,output)
        elif operation == 'flatten':
            if pages != list(range(len(doc))): raise ValueError('Achatar PDF aplica-se ao documento inteiro; remova a seleção de páginas.')
            annotations=_boolean(options.get('annotations',True)); widgets=_boolean(options.get('widgets',True))
            if not annotations and not widgets: raise ValueError('Selecione anotações ou formulários para achatar.')
            doc.bake(annots=annotations,widgets=widgets); save(doc,output)
        elif operation in ('delete-pages','extract-pages'):
            if not options.get('pages'): raise ValueError('Informe as páginas que deseja excluir ou extrair.')
            omitted=set(pages)
            selected=[p for p in range(len(doc)) if p not in omitted] if operation=='delete-pages' else pages
            if not selected: raise ValueError('Não é possível excluir todas as páginas.')
            doc.select(selected); save(doc,output)
        elif operation == 'organize':
            if not options.get('pages'): raise ValueError('Informe a nova ordem das páginas.')
            doc.select(pages)
            secret=options.get('output_password')
            if secret is not None:
                if not isinstance(secret,str) or not 8<=len(secret)<=40: raise ValueError('A senha do resultado deve conter de 8 a 40 caracteres.')
                save(doc,output,encryption=fitz.PDF_ENCRYPT_AES_256,user_pw=secret,owner_pw=secrets.token_urlsafe(24),permissions=doc.permissions)
            else: save(doc,output)
        elif operation == 'form-create':
            _create_forms(doc,options); save(doc,output)
        elif operation == 'form-fill':
            _fill_forms(doc,options); save(doc,output)
        elif operation in PDF_OPERATIONS:
            _modify(doc,pages,operation,options)
            kwargs = {}
            if operation == 'protect':
                secret = options.get('new_password','')
                if not isinstance(secret,str) or not 8 <= len(secret) <= 40: raise ValueError('Use uma senha de 8 a 40 caracteres.')
                kwargs = dict(encryption=fitz.PDF_ENCRYPT_AES_256,owner_pw=secrets.token_urlsafe(24),user_pw=secret)
            elif operation == 'unlock': kwargs['encryption'] = fitz.PDF_ENCRYPT_NONE
            save(doc,output,**kwargs)
        else: raise ValueError('Operação não implementada: '+operation)


def _paper_geometry(options):
    width,height=fitz.paper_size(str(options.get('paper','a4')).lower())
    if width<=0 or height<=0: raise ValueError('Tamanho de papel desconhecido.')
    if options.get('landscape') or options.get('orientation')=='landscape': width,height=height,width
    margin=number(options.get('margin',36),'Margem',0,min(width,height)/3)
    return width,height,margin


def _text_file_pdf(source,output,options):
    encoding=str(options.get('encoding','utf-8-sig'))
    try: text=Path(source).read_text(encoding=encoding,errors='strict')
    except (UnicodeError,LookupError) as exc:
        raise ValueError(f'Não foi possível ler o texto como {encoding}. Selecione a codificação correta.') from exc
    width,height,margin=_paper_geometry(options)
    fontsize=number(options.get('font_size',11),'Tamanho da fonte',4,72)
    # Story ignores overflow-wrap for unbroken tokens. Invisible break opportunities
    # preserve ASCII token content while preventing URLs/identifiers from clipping.
    measure=fitz.Font(fontname='helv'); available=width-2*margin
    def soft_break(match):
        token=match.group()
        if token.isascii() and measure.text_length(token,fontsize=fontsize)>available*0.85:
            return '\u200b'.join(token)
        return token
    flow_text=re.sub(r'\S+',soft_break,text)
    # Escape source HTML and preserve line boundaries while allowing long lines to wrap.
    content='<html><body>'+''.join('<p>'+html.escape(line).replace('\t','&nbsp;&nbsp;&nbsp;&nbsp;')+'</p>' if line else '<p>&nbsp;</p>' for line in flow_text.splitlines())+'</body></html>'
    if not text.strip(): content='<html><body><p>&nbsp;</p></body></html>'
    story=fitz.Story(content,user_css=f'body {{font-family:sans-serif;font-size:{fontsize}pt}} p {{margin:0 0 0.3em 0;white-space:pre-wrap;overflow-wrap:anywhere}}')
    bounds=fitz.Rect(0,0,width,height); area=bounds+(margin,margin,-margin,-margin)
    def layout(rect_number,filled):
        if rect_number>=10000: raise ValueError('O texto excede o limite de 10.000 páginas.')
        return bounds,area,None
    writer=fitz.DocumentWriter(str(output))
    try: story.write(writer,layout)
    finally: writer.close()
    # Refuse unsupported over-wide Unicode runs instead of publishing cropped text.
    with fitz.open(output) as rendered:
        for page in rendered:
            bounds=page.rect+(-2,-2,2,2)
            for block in page.get_text('dict',clip=fitz.INFINITE_RECT())['blocks']:
                if block['type']==0 and not bounds.contains(fitz.Rect(block['bbox'])):
                    raise ValueError('Um trecho de texto não cabe na página. Insira quebras de linha, diminua a fonte ou use papel maior; nenhuma cópia foi criada.')


def _csv_workbook(source,output,options):
    import csv
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter
    import textwrap
    import unicodedata
    encoding=str(options.get('encoding','utf-8-sig'))
    delimiter=options.get('delimiter')
    if delimiter is not None and (not isinstance(delimiter,str) or len(delimiter)!=1 or delimiter in ('\n','\r','"')):
        raise ValueError('O separador CSV deve ser um único caractere válido.')
    workbook=Workbook(); sheet=workbook.active; sheet.title='Dados'
    widths={}
    maximum_width=number(options.get('max_column_width',60),'Largura máxima de coluna',12,120)
    def display_width(value):
        return sum(2 if unicodedata.east_asian_width(character) in ('W','F') else 1 for character in value)
    try:
        with Path(source).open('r',encoding=encoding,errors='strict',newline='') as handle:
            sample=handle.read(8192); handle.seek(0)
            if delimiter is None:
                try: dialect=csv.Sniffer().sniff(sample,delimiters=',;\t|')
                except csv.Error: dialect=csv.excel
                reader=csv.reader(handle,dialect)
            else: reader=csv.reader(handle,delimiter=delimiter)
            for row_number,row in enumerate(reader,1):
                if row_number>1048576 or len(row)>16384: raise ValueError('O CSV excede os limites de linhas ou colunas da planilha.')
                sheet.append(row)
                for cell in sheet[row_number]:
                    cell.data_type='s'  # Content never becomes a spreadsheet formula.
                    cell.alignment=Alignment(wrap_text=True,vertical='top')
                    longest=max((display_width(line) for line in str(cell.value or '').splitlines()),default=0)
                    widths[cell.column]=max(widths.get(cell.column,12),min(maximum_width,longest+3))
    except (UnicodeError,LookupError) as exc:
        raise ValueError(f'Não foi possível ler o CSV como {encoding}. Selecione a codificação correta.') from exc
    for column,width in widths.items(): sheet.column_dimensions[get_column_letter(column)].width=width
    for row in sheet.iter_rows():
        lines=1
        for cell in row:
            available=max(1,int(widths.get(cell.column,12))-3)
            required=0
            for line in str(cell.value or '').split('\n'):
                # Keep generous height for wrapped words and double-width Unicode glyphs.
                wrapped=textwrap.wrap(line,width=available,replace_whitespace=False,drop_whitespace=False) or ['']
                required+=sum(max(1,math.ceil(display_width(part)/available)) for part in wrapped)
            lines=max(lines,required)
        height=lines*18+8
        if height>400:
            workbook.close()
            raise ValueError('Uma célula CSV exige mais altura do que a planilha pode exibir. Divida o conteúdo da célula ou aumente max_column_width; nenhuma cópia foi criada.')
        sheet.row_dimensions[row[0].row].height=height
    try: workbook.save(output)
    finally: workbook.close()
    return output


def _raster_export(doc,pages,output,format_name,options):
    output.mkdir(); dpi=integer(options.get('dpi',150),'DPI',36,600)
    if options.get('mode')=='embedded':
        seen=set(); count=0
        for i in pages:
            for item in doc[i].get_images(full=True):
                if item[0] in seen: continue
                seen.add(item[0]); count+=1
                if format_name=='png':
                    pix=fitz.Pixmap(doc,item[0])
                    if pix.colorspace and pix.colorspace.n>3: pix=fitz.Pixmap(fitz.csRGB,pix)
                    if item[1] and not pix.alpha: pix=fitz.Pixmap(pix,fitz.Pixmap(doc,item[1]))
                    pix.save(str(output/f'imagem-{count:04}.png'))
                else:
                    from PIL import Image
                    data=doc.extract_image(item[0])
                    with Image.open(io.BytesIO(data['image'])) as image:
                        image.convert('RGB').save(output/f'imagem-{count:04}.jpg',quality=95)
        if count==0: raise ValueError('Nenhuma imagem incorporada foi encontrada.')
    else:
        alpha=_boolean(options.get('alpha',False)) if format_name=='png' else False
        for i in pages:
            doc[i].get_pixmap(dpi=dpi,alpha=alpha).save(str(output/f'pagina-{i+1:04}.{format_name}'))


def _font(options):
    font_file = options.get('font_file')
    if font_file:
        path = Path(font_file)
        if not path.is_file(): raise ValueError('Arquivo de fonte não encontrado.')
        return 'customfont',str(path)
    font = options.get('font','helv')
    if font not in ('helv','hebo','heit','hebi','tiro','tibo','tiit','tibi','cour','cobo','coit','cobi'):
        raise ValueError('Fonte desconhecida. Escolha Helvetica, Times, Courier ou font_file.')
    return font,None


def _checked_text(value):
    if not isinstance(value,str): raise ValueError('O texto deve ser uma string.')
    if any(ord(character)<32 and character not in '\n\r\t' for character in value):
        raise ValueError('O texto contém caracteres de controle não suportados.')
    return value


def _font_has_text(text,font,fontfile):
    # Standard PDF insertion uses a single-byte encoding even when its Font
    # object has extra glyphs. HTML shaping provides Unicode fallback instead.
    if not fontfile: return all(ord(character)<=255 for character in text)
    measure=fitz.Font(fontfile=fontfile)
    if any(not measure.has_glyph(ord(character)) for character in text if not character.isspace()):
        raise ValueError('A fonte escolhida não contém todos os caracteres. Escolha outra fonte; nenhuma cópia foi criada.')
    return True


def _html_text(page,rect,text,fontsize,ink,minimum=None,opacity=1,align=0,font='helv'):
    """Shape on a disposable page and verify bounds before importing the text."""
    minimum=fontsize if minimum is None else minimum
    family='monospace' if font.startswith('co') else 'serif' if font.startswith('ti') else 'sans-serif'
    weight='bold' if font in ('hebo','hebi','tibo','tibi','cobo','cobi') else 'normal'
    style='italic' if font in ('heit','hebi','tiit','tibi','coit','cobi') else 'normal'
    css_color='#'+''.join(f'{round(component*255):02x}' for component in ink)
    content='<div>'+html.escape(text).replace('\n','<br>').replace('\t','&nbsp;&nbsp;&nbsp;&nbsp;')+'</div>'
    css=f'* {{font-family:{family};font-size:{fontsize}pt;color:{css_color};font-weight:{weight};font-style:{style};margin:0;padding:0;line-height:1.5;text-align:{("left","center","right","justify")[align]}}}'
    with fitz.open() as shaped:
        scratch=shaped.new_page(width=max(1,rect.x1+2),height=max(1,rect.y1+2))
        spare,scale=scratch.insert_htmlbox(rect,content,css=css,scale_low=minimum/fontsize,opacity=opacity)
        if spare<0: raise ValueError('O texto não cabe na área escolhida. Reduza o texto ou a fonte; nenhuma cópia foi criada.')
        for trace in scratch.get_texttrace():
            if any(glyph==0 and not chr(codepoint).isspace() for codepoint,glyph,*_ in trace['chars']):
                raise ValueError('As fontes disponíveis não contêm todos os caracteres; nenhuma cópia foi criada.')
        visible=[]
        for block in scratch.get_text('dict',clip=fitz.INFINITE_RECT())['blocks']:
            if block['type']!=0: continue
            for line in block['lines']:
                for span in line['spans']:
                    if span['text'].strip(): visible.append(fitz.Rect(span['bbox']))
        # MuPDF may report a fit for over-wide unbroken words; reject clipping.
        tolerance=rect+(-1,-1,1,1)
        if text.strip() and (not visible or any(not tolerance.contains(bounds) for bounds in visible)):
            raise ValueError('O texto excede a área disponível; nenhuma cópia foi criada.')
        if text.strip(): page.show_pdf_page(rect,shaped,0,clip=rect,keep_proportion=False)


def _text(page,text,options,operation):
    text=_checked_text(text)
    bounds = _page_bounds(page)
    fontsize = number(options.get('font_size',14),'Tamanho da fonte',4,300)
    opacity = number(options.get('opacity',0.5 if operation=='watermark' else 1),'Opacidade',0,1)
    ink = color(options.get('color'))
    font,fontfile = _font(options)
    standard=_font_has_text(text,font,fontfile)
    margin = number(options.get('margin',36),'Margem',0,min(bounds.width,bounds.height)/2)
    if 'rect' in options or 'width' in options or 'height' in options:
        rect = rectangle(options,page)
        alignment=integer(options.get('align',0),'Alinhamento',0,3)
        if not standard:
            _html_text(page,rect,text,fontsize,ink,opacity=opacity,align=alignment,font=font); return
        remaining = page.insert_textbox(rect,text,fontname=font,fontfile=fontfile,fontsize=fontsize,color=ink,fill_opacity=opacity,align=alignment)
        if remaining < 0: raise ValueError('O texto não cabe na área escolhida. Aumente a área ou diminua a fonte.')
    else:
        measure = fitz.Font(fontfile=fontfile) if fontfile else fitz.Font(fontname=font)
        width = max((measure.text_length(line,fontsize=fontsize) for line in text.splitlines()),default=0)
        position = options.get('position','bottom-left' if operation=='numbers' else 'top-left')
        if position not in ('top-left','top-center','top-right','center','bottom-left','bottom-center','bottom-right'):
            raise ValueError('Posição desconhecida.')
        x = margin if position.endswith('left') else (bounds.width-width)/2 if position.endswith('center') or position=='center' else bounds.width-margin-width
        y = bounds.height-margin if position.startswith('bottom') else bounds.height/2 if position=='center' else margin+fontsize
        x = number(options.get('x',x),'x'); y = number(options.get('y',y),'y')
        if x < 0 or y-fontsize < 0 or x+width > bounds.width or y > bounds.height:
            raise ValueError('O texto não cabe na página. Ajuste posição ou tamanho.')
        if not standard:
            rect=fitz.Rect(x,y-fontsize,min(bounds.width,x+max(width+fontsize,fontsize)),min(bounds.height,y+fontsize*max(1,len(text.splitlines()))*1.6))
            _html_text(page,rect,text,fontsize,ink,opacity=opacity,font=font)
        else:
            # Measure the actual rendered spans, including descenders and multiline
            # baselines. insert_text silently drops lines that cross page bounds.
            with fitz.open() as measured:
                scratch=measured.new_page(width=bounds.width,height=max(bounds.height,y+fontsize*max(2,len(text.splitlines()))*3))
                scratch.insert_text((x,y),text,fontname=font,fontfile=fontfile,fontsize=fontsize,color=ink,fill_opacity=opacity)
                for block in scratch.get_text('dict',clip=fitz.INFINITE_RECT())['blocks']:
                    if block['type']==0 and not bounds.contains(fitz.Rect(block['bbox'])):
                        raise ValueError('O texto não cabe na página. Ajuste posição ou tamanho; nenhuma cópia foi criada.')
            page.insert_text((x,y),text,fontname=font,fontfile=fontfile,fontsize=fontsize,color=ink,fill_opacity=opacity)


def _reject_pending_redactions(page):
    if any(annotation.type[0]==fitz.PDF_ANNOT_REDACT for annotation in (page.annots() or [])):
        raise ValueError('Esta página contém marcações de redação pendentes. Resolva essas marcações antes de editar; nenhuma cópia foi criada.')


def _remove_redacted_interactive(page,rects,term):
    def overlaps(bounds): return any(_area(fitz.Rect(bounds)&rect)>0 for rect in rects)
    def sensitive(value): return bool(term) and str(term).casefold() in str(value or '').casefold()
    for widget in list(page.widgets() or []):
        if overlaps(widget.rect) or sensitive(widget.field_value) or sensitive(widget.field_label):
            page.delete_widget(widget)
    for annotation in list(page.annots() or []):
        if overlaps(annotation.rect) or any(sensitive(value) for value in annotation.info.values()):
            page.delete_annot(annotation)
    for link in page.get_links():
        if overlaps(link['from']) or any(sensitive(link.get(key)) for key in ('uri','file','name')):
            page.delete_link(link)


def _modify(doc,pages,operation,options):
    if operation == 'compress' and options.get('image_dpi'):
        # Recompress embedded images while preserving vectors, searchable text and widgets.
        from PIL import Image
        dpi = number(options['image_dpi'],'DPI',36,600)
        quality = int(number(options.get('quality',80),'Qualidade JPEG',20,95))
        placements={}; selected=set(pages)
        for index,page in enumerate(doc):
            for item in page.get_images(full=True):
                xref=item[0]
                record=placements.setdefault(xref,dict(page=index,width=0,height=0,masked=False,eligible=True))
                record['masked'] |= bool(item[1])
                record['eligible'] &= index in selected
                for rect in page.get_image_rects(xref):
                    record['width']=max(record['width'],rect.width)
                    record['height']=max(record['height'],rect.height)
        for xref,record in placements.items():
            if record['masked'] or not record['eligible'] or not record['width'] or not record['height']: continue
            target_w=max(1,math.ceil(record['width']*dpi/72)); target_h=max(1,math.ceil(record['height']*dpi/72))
            try:
                source=doc.extract_image(xref)
                with Image.open(io.BytesIO(source['image'])) as image:
                    # A shared image must retain enough pixels for its largest placement.
                    ratio=max(target_w/image.width,target_h/image.height)
                    if ratio>=1: continue
                    image.thumbnail((math.ceil(image.width*ratio),math.ceil(image.height*ratio)))
                    buffer=io.BytesIO(); image.convert('RGB').save(buffer,format='JPEG',quality=quality,optimize=True)
                    doc[record['page']].replace_image(xref,stream=buffer.getvalue())
            except (OSError,ValueError): continue
    matches = 0
    for i in pages:
        page = doc[i]
        if operation == 'rotate':
            degrees = number(options.get('degrees',90),'Rotação')
            if degrees % 90: raise ValueError('A rotação deve ser múltipla de 90 graus.')
            page.set_rotation((page.rotation+int(degrees))%360)
        elif operation in ('text','watermark','numbers'):
            if operation == 'watermark' and options.get('image'):
                _insert_image(page,options,watermark=True)
            else:
                value = options.get('text','PDF')
                if operation == 'numbers':
                    start = integer(options.get('start',1),'Número inicial',0)
                    value = options.get('format','{n}').replace('{n}',str(start+pages.index(i))).replace('{total}',str(len(pages)))
                _text(page,str(value),options,operation)
        elif operation in ('image','signature'): _insert_image(page,options)
        elif operation == 'shape':
            rect = rectangle(options,page)
            drawing = page.new_shape(); kind = options.get('shape','rectangle')
            if kind == 'rectangle': drawing.draw_rect(rect)
            elif kind == 'ellipse': drawing.draw_oval(rect)
            elif kind == 'line': drawing.draw_line(rect.tl,rect.br)
            else: raise ValueError('Escolha rectangle, ellipse ou line.')
            drawing.finish(color=color(options.get('color')),fill=color(options['fill']) if options.get('fill') else None,width=number(options.get('line_width',1),'Espessura',0.1,50),fill_opacity=number(options.get('opacity',1),'Opacidade',0,1),stroke_opacity=number(options.get('opacity',1),'Opacidade',0,1))
            drawing.commit()
        elif operation in ('redact','edit'):
            term = options.get('search',options.get('text',''))
            rects = [rectangle({'rect':r},page) for r in options.get('rects',[])] if operation=='redact' and options.get('rects') else page.search_for(term) if term else []
            if not term and not rects: raise ValueError('Informe o texto ou as áreas a remover.')
            replacement = _checked_text(options.get('replacement',''))
            font,_ = _font(options)
            if operation=='edit' and options.get('font_file'): raise ValueError('Substituição direta usa fontes PDF padrão; use sobreposição para fontes externas.')
            if rects:
                _reject_pending_redactions(page)
                if operation=='redact': _remove_redacted_interactive(page,rects,term)
            for rect in rects:
                matches += 1
                page.add_redact_annot(rect,fill=color(options.get('fill'),(1,1,1) if operation=='edit' else (0,0,0)))
            if rects:
                if operation=='edit': page.apply_redactions(images=0,graphics=0,text=0)
                else: page.apply_redactions()
            if operation=='edit' and replacement:
                for rect in rects:
                    size=number(options.get('font_size',11),'Fonte',4,300)
                    minimum=number(options.get('min_font_size',min(8,size)),'Fonte mínima',4,size)
                    if not _font_has_text(replacement,font,None):
                        _html_text(page,rect,replacement,size,color(options.get('color')),minimum=minimum,font=font)
                        continue
                    fitted=False
                    while size >= minimum:
                        spare=page.insert_textbox(rect,str(replacement),fontname=font,fontsize=size,color=color(options.get('color')))
                        if spare>=0: fitted=True; break
                        size-=0.5
                    if not fitted: raise ValueError('O texto substituto não cabe na área original. Reduza o texto ou ajuste a fonte mínima; nenhuma cópia foi criada.')
        elif operation == 'crop':
            margin = number(options.get('margin',20),'Margem',0)
            rect = rectangle(options,page)+tuple(page.cropbox.tl)*2 if 'rect' in options else page.cropbox+(margin,margin,-margin,-margin)
            if rect.is_empty: raise ValueError('Margem inválida.')
            page.set_cropbox(rect)
    if operation in ('redact','edit') and matches == 0:
        raise ValueError('Nenhuma ocorrência encontrada. Nenhuma cópia foi criada.')
    if operation == 'redact':
        # Purge metadata, embedded attachments and XML metadata along with removed streams.
        doc.set_metadata({})
        doc.del_xml_metadata()
        for name in list(doc.embfile_names()): doc.embfile_del(name)
        doc.scrub(attached_files=False,clean_pages=False,embedded_files=False,hidden_text=False,javascript=True,metadata=False,redactions=False,remove_links=False,reset_fields=False,reset_responses=False,thumbnails=True,xml_metadata=False)
        term=options.get('search',options.get('text',''))
        if term:
            outline=doc.get_toc()
            for item in outline:
                if str(term).casefold() in item[1].casefold(): item[1]='[Removido]'
            if outline: doc.set_toc(outline)



def _insert_image(page,options,watermark=False):
    image = options.get('image')
    if not image or not Path(image).is_file(): raise ValueError('Selecione uma imagem válida.')
    rect = rectangle(options,page)
    opacity = number(options.get('opacity',0.35 if watermark else 1),'Opacidade',0,1)
    if opacity < 1:
        from PIL import Image
        with Image.open(image) as source:
            rgba = source.convert('RGBA'); rgba.putalpha(rgba.getchannel('A').point(lambda a:int(a*opacity)))
            buffer = io.BytesIO(); rgba.save(buffer,format='PNG')
            page.insert_image(rect,stream=buffer.getvalue(),keep_proportion=True,overlay=bool(options.get('overlay',True)))
    else: page.insert_image(rect,filename=image,keep_proportion=True,overlay=bool(options.get('overlay',True)))


def _layout(doc,pages):
    blocks=[]
    for i in pages:
        for j,block in enumerate(doc[i].get_text('dict',sort=True)['blocks']):
            if block['type']!=0: continue
            text='\n'.join(''.join(span['text'] for span in line['spans']) for line in block['lines']).strip()
            if not text: continue
            spans=[span for line in block['lines'] for span in line['spans'] if span['text'].strip()]
            if not spans: continue
            ink=spans[0].get('color',0)
            blocks.append(dict(id=f'p{i+1}-b{j}',page=i+1,rect=list(block['bbox']),text=text,font_size=max(span['size'] for span in spans),color=f'#{ink:06x}'))
    return blocks


def _translate_layout(doc,options):
    edits=options.get('edits',[])
    if not isinstance(edits,list) or not edits: raise ValueError('Informe os blocos traduzidos.')
    grouped={}
    for edit in edits:
        page_number=integer(edit.get('page',1),'Página',1,len(doc))-1
        page=doc[page_number]; rect=rectangle(edit,page)
        text=str(edit.get('text',''))
        if not text.strip(): raise ValueError('O texto traduzido não pode estar vazio.')
        if edit.get('source_text') and page.get_textbox(rect).strip()!=edit['source_text'].strip():
            raise ValueError('O bloco original mudou. Extraia o layout novamente.')
        if any(_area(rect & previous[1]) > 0.5 for previous in grouped.get(page_number,[])):
            raise ValueError('Blocos de tradução sobrepostos. Selecione blocos sem sobreposição.')
        grouped.setdefault(page_number,[]).append((edit,rect,text))
    for index,items in grouped.items():
        page=doc[index]
        _reject_pending_redactions(page)
        for edit,rect,text in items: page.add_redact_annot(rect,fill=False)
        page.apply_redactions(images=0,graphics=0,text=0)
        for edit,rect,text in items:
            size=number(edit.get('font_size',11),'Fonte',4,300)
            minimum=number(edit.get('min_font_size',options.get('min_font_size',min(8,size))),'Fonte mínima',4,size)
            _html_text(page,rect,_checked_text(text),size,color(edit.get('color'),(0,0,0)),minimum=minimum)


def _detect_forms(doc,pages):
    """Suggest vector boxes/blank lines for human review, never claim OCR inference."""
    fields=[]
    for i in pages:
        page=doc[i]; rectangles=[]
        for drawing in page.get_drawings():
            for item in drawing['items']:
                if item[0]=='re':
                    rect=fitz.Rect(item[1])
                    if 8<=rect.width<=400 and 8<=rect.height<=100: rectangles.append(rect)
                elif item[0]=='l':
                    start,end=item[1:3]
                    if abs(start.y-end.y)<1 and abs(start.x-end.x)>=35:
                        rectangles.append(fitz.Rect(min(start.x,end.x),max(0,start.y-18),max(start.x,end.x),start.y+2))
        seen=[]
        for rect in rectangles:
            if not _page_bounds(page).contains(rect) or any(_area(rect & previous)>_area(rect)*0.8 for previous in seen): continue
            # Occupied cells are labels/content, not blank input fields.
            if page.get_textbox(rect).strip(): continue
            seen.append(rect)
            kind='checkbox' if rect.width<=24 and rect.height<=24 else 'text'
            label=page.get_textbox(fitz.Rect(max(0,rect.x0-180),max(0,rect.y0-20),rect.x1,rect.y0)).strip()
            fields.append(dict(name=f'campo_{i+1}_{len(seen)}',page=i+1,type=kind,rect=list(rect),label=label[-120:],confidence='heuristic'))
    return fields


def _pdf_name(value):
    return '/'+''.join(chr(byte) if 33<=byte<=126 and chr(byte) not in '#%()/<>[]{}' else f'#{byte:02X}' for byte in value.encode('utf-8'))


def _radio_value(widget):
    value=widget.on_state()
    if not isinstance(value,str): raise ValueError('O botão de escolha não possui um estado válido.')
    data=re.sub(r'#([0-9a-fA-F]{2})',lambda match:chr(int(match[1],16)),value)
    try: return data.encode('latin1').decode('utf-8')
    except (UnicodeError,ValueError): return data


def _form_records(doc,pages):
    records=[]; groups={}
    for index in pages:
        for widget in doc[index].widgets() or []:
            record=dict(page=index+1,name=widget.field_name,value=widget.field_value,type=widget.field_type_string,rect=list(widget.rect),choices=widget.choice_values,readonly=bool(widget.field_flags & 1))
            if widget.field_type!=fitz.PDF_WIDGET_TYPE_RADIOBUTTON:
                records.append(record); continue
            key=widget.rb_parent or widget.field_name
            if key not in groups:
                record.update(value='',choices=[],buttons=[]); groups[key]=record; records.append(record)
            group=groups[key]; group['readonly'] |= bool(widget.field_flags & 1); value=_radio_value(widget)
            group['choices'].append(value); group['buttons'].append(dict(page=index+1,rect=list(widget.rect),value=value))
            if widget.field_value!='Off': group['value']=value
    return records


def _create_radio(doc,field,name):
    choices=field.get('choices',[])
    if not isinstance(choices,list) or len(choices)<2: raise ValueError('Um grupo de escolha precisa de pelo menos duas opções com value e rect.')
    values=[]; widgets=[]
    for choice in choices:
        if not isinstance(choice,dict): raise ValueError('Cada opção precisa de value, rect e, opcionalmente, page.')
        value=choice.get('value')
        if not isinstance(value,str) or not value or value=='Off' or value in values: raise ValueError('As opções devem ter valores únicos, não vazios e diferentes de Off.')
        values.append(value)
        index=integer(choice.get('page',field.get('page',1)),'Página',1,len(doc))-1
        page=doc[index]; widget=fitz.Widget(); widget.field_name=name+'_'+str(len(values)); widget.field_type=fitz.PDF_WIDGET_TYPE_CHECKBOX
        widget.rect=rectangle(choice,page); widget.field_value='Off'; widget.border_width=1; widget.border_color=color(field.get('border_color'),(.5,.5,.5))
        widgets.append(page.add_widget(widget).xref)
        if choice.get('label') is not None:
            label=_checked_text(choice['label']); rect=widget.rect
            label_rect=choice.get('label_rect',[rect.x1+6,rect.y0,_page_bounds(page).x1-12,max(rect.y1,rect.y0+24)])
            _text(page,label,dict(rect=label_rect,font_size=field.get('font_size',12),color=field.get('color')),'text')
    selected=field.get('value','')
    if selected not in ['',*values]: raise ValueError('O valor inicial não pertence às opções do grupo.')
    parent=doc.get_new_xref()
    doc.update_object(parent,'<</FT/Btn/Ff 32768/T '+fitz.get_pdf_str(name)+'/V '+(_pdf_name(selected) if selected else '/Off')+'/Kids['+' '.join(f'{xref} 0 R' for xref in widgets)+']>>')
    for xref,value in zip(widgets,values):
        off=doc.xref_get_key(xref,'AP/N/Off')[1]; on=doc.xref_get_key(xref,'AP/N/Yes')[1]
        doc.xref_set_key(xref,'AP/N',f'<</Off {off} {_pdf_name(value)} {on}>>')
        for key in ('FT','Ff','T','V'): doc.xref_set_key(xref,key,'null')
        doc.xref_set_key(xref,'Parent',f'{parent} 0 R'); doc.xref_set_key(xref,'AS',_pdf_name(value) if value==selected else '/Off')
    # Each group is one terminal field; its widget annotations remain on their pages.
    catalog=doc.pdf_catalog(); roots=doc.xref_get_key(catalog,'AcroForm/Fields')[1]
    references=[int(match[1]) for match in re.finditer(r'(\d+)\s+0\s+R',roots)]
    references=[xref for xref in references if xref not in widgets]+[parent]
    doc.xref_set_key(catalog,'AcroForm/Fields','['+' '.join(f'{xref} 0 R' for xref in references)+']')


def _create_forms(doc,options):
    fields=options.get('fields',[])
    if not isinstance(fields,list) or not fields: raise ValueError('Defina ao menos um campo.')
    names={widget.field_name for page in doc for widget in (page.widgets() or [])}
    types={'text':fitz.PDF_WIDGET_TYPE_TEXT,'checkbox':fitz.PDF_WIDGET_TYPE_CHECKBOX,'choice':fitz.PDF_WIDGET_TYPE_COMBOBOX,'list':fitz.PDF_WIDGET_TYPE_LISTBOX}
    for field in fields:
        if not isinstance(field,dict): raise ValueError('Cada campo deve ser um objeto JSON.')
        name=field.get('name','')
        if not isinstance(name,str) or not name.strip() or name in names: raise ValueError('Os campos precisam de nomes únicos e não vazios.')
        names.add(name); kind=field.get('type','text')
        if kind=='radio': _create_radio(doc,field,name); continue
        if kind not in types: raise ValueError('Tipo de campo inválido.')
        index=integer(field.get('page',1),'Página',1,len(doc))-1
        page=doc[index]; widget=fitz.Widget(); widget.field_name=name; widget.field_type=types[kind]
        widget.rect=rectangle(field,page); widget.text_fontsize=number(field.get('font_size',12),'Fonte',4,100)
        widget.text_color=color(field.get('color'),(0,0,0)); widget.border_color=color(field.get('border_color'),(.5,.5,.5)); widget.border_width=1
        if kind in ('choice','list'):
            choices=field.get('choices',field.get('options',[]))
            if not isinstance(choices,list) or not choices or any(not isinstance(value,str) or not value for value in choices) or len(set(choices))!=len(choices):
                raise ValueError('Defina uma lista de opções únicas e não vazias.')
            if any(not _font_has_text(value,'helv',None) for value in choices): raise ValueError('A aparência de formulário exige caracteres Latin-1; use uma fonte de formulário compatível em outro editor para outros alfabetos.')
            widget.choice_values=choices
        value=field.get('value','')
        if kind=='checkbox': widget.field_value='Yes' if _boolean(value) else 'Off'
        else:
            value=_checked_text(value)
            if kind in ('choice','list') and value and value not in choices: raise ValueError('O valor inicial não pertence às opções do campo.')
            if not _font_has_text(value,'helv',None): raise ValueError('A aparência de formulário exige caracteres Latin-1; nenhuma cópia foi criada.')
            widget.field_value=value
        if _boolean(field.get('multiline',False)) and kind=='text': widget.field_flags=fitz.PDF_TX_FIELD_IS_MULTILINE
        page.add_widget(widget)


def _fill_forms(doc,options):
    values=options.get('values',{})
    if not isinstance(values,dict) or not values or any(not isinstance(name,str) for name in values): raise ValueError('Informe os valores dos campos por nome.')
    existing={widget.field_name for page in doc for widget in (page.widgets() or [])}
    missing=set(values)-existing
    if missing: raise ValueError('Campos inexistentes: '+', '.join(sorted(missing)))
    radio={}
    for page in doc:
        for widget in page.widgets() or []:
            if widget.field_name not in values: continue
            if widget.field_flags & 1: raise ValueError('O campo é somente leitura: '+widget.field_name)
            value=values[widget.field_name]
            if widget.field_type==fitz.PDF_WIDGET_TYPE_RADIOBUTTON:
                radio.setdefault(widget.field_name,[]).append((widget.xref,widget.rb_parent,_radio_value(widget))); continue
            if widget.field_type==fitz.PDF_WIDGET_TYPE_CHECKBOX: widget.field_value=widget.on_state() if _boolean(value) else 'Off'
            elif widget.field_type in (fitz.PDF_WIDGET_TYPE_TEXT,fitz.PDF_WIDGET_TYPE_LISTBOX,fitz.PDF_WIDGET_TYPE_COMBOBOX):
                value=_checked_text(value)
                if widget.field_type!=fitz.PDF_WIDGET_TYPE_TEXT and value and value not in (widget.choice_values or []): raise ValueError('Escolha inválida para '+widget.field_name)
                if not _font_has_text(value,'helv',None): raise ValueError('A aparência de formulário exige caracteres Latin-1; nenhuma cópia foi criada.')
                widget.field_value=value
                if value=='':
                    doc.xref_set_key(widget.xref,'V','()')
                    if widget.field_type!=fitz.PDF_WIDGET_TYPE_TEXT: doc.xref_set_key(widget.xref,'I','[]')
            else: raise ValueError('Este tipo de campo não pode ser preenchido como texto: '+widget.field_name)
            widget.update()
    for name,buttons in radio.items():
        value=values[name]
        if not isinstance(value,str) or value not in ['',*[item[2] for item in buttons]]: raise ValueError('Escolha inválida para '+name)
        for xref,parent,choice in buttons:
            doc.xref_set_key(xref,'AS',_pdf_name(choice) if choice==value else '/Off')
            doc.xref_set_key(parent or xref,'V',_pdf_name(value) if value else '/Off')


def _page_elements(page):
    """Reading-order blocks and extracted tables, avoiding duplicated table text."""
    tables = page.find_tables().tables
    entries = [dict(type=2,bbox=table.bbox,rows=table.extract()) for table in tables]
    for block in page.get_text('dict',sort=True)['blocks']:
        bounds = fitz.Rect(block['bbox'])
        if block['type']==0 and any(_area(bounds & fitz.Rect(table.bbox)) > _area(bounds)*0.65 for table in tables):
            continue
        entries.append(block)
    return sorted(entries,key=lambda b:(round(b['bbox'][1]/8),b['bbox'][0]))


def _write_base64(stream, data):
    # A multiple of three preserves the encoding across chunks without padding.
    with memoryview(data) as view:
        for offset in range(0,len(view),49152):
            stream.write(base64.b64encode(view[offset:offset+49152]).decode('ascii'))


def _markdown(doc,pages,output):
    def escaped(value):
        return str(value or '').replace('\\','\\\\').replace('<','&lt;').replace('>','&gt;')
    with output.open('w',encoding='utf-8') as stream:
        first=True
        def begin_line():
            nonlocal first
            if not first: stream.write('\n')
            first=False
        def emit(text):
            begin_line(); stream.write(text)
        for i in pages:
            page=doc[i]; elements=_page_elements(page)
            sizes=[s['size'] for b in elements if b['type']==0 for line in b['lines'] for s in line['spans'] if s['text'].strip()]
            body=sorted(sizes)[len(sizes)//2] if sizes else 11
            emit(f'<!-- Página {i+1} -->')
            for block in elements:
                if block['type']==2:
                    rows=block['rows']
                    if not rows: continue
                    width=max(map(len,rows))
                    for n,row in enumerate(rows):
                        emit('| '+' | '.join(escaped(value).replace('|','\\|').replace('\n','<br>') for value in [*row,*['']*(width-len(row))])+' |')
                        if n==0: emit('| '+' | '.join(['---']*width)+' |')
                    emit('')
                    continue
                if block['type']!=0:
                    # Stream image encoding instead of retaining every page's data URI.
                    begin_line(); stream.write(f'![Imagem da página {i+1}](data:image/{block.get("ext","png")};base64,')
                    _write_base64(stream,block['image']); stream.write(')')
                    continue
                for line in block['lines']:
                    text=''.join(span['text'] for span in line['spans']).strip()
                    if not text: continue
                    size=max(span['size'] for span in line['spans'])
                    prefix='# ' if size>=body*1.7 else '## ' if size>=body*1.3 else ''
                    emit(prefix+escaped(text))
                emit('')
            for link in page.get_links():
                if link.get('uri'):
                    uri=str(link['uri'])
                    if uri.startswith(('https://','http://','mailto:')):
                        label=escaped(page.get_textbox(link['from']).strip() or 'Link').replace('[','\\[').replace(']','\\]')
                        emit(f'[{label}]({uri.replace("(","%28").replace(")","%29")})')


def _word(doc,pages,output):
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    result = Document()
    for count,i in enumerate(pages):
        page = doc[i]
        if count: result.add_page_break()
        for block in _page_elements(page):
            if block['type']==2:
                rows=block['rows']
                if rows:
                    table=result.add_table(rows=len(rows),cols=max(map(len,rows)))
                    table.style='Table Grid'
                    for r,row in enumerate(rows):
                        for c,value in enumerate(row): table.cell(r,c).text=str(value or '')
                continue
            if block['type']==1:
                try: result.add_picture(io.BytesIO(block['image']),width=Inches(min(6,(block['bbox'][2]-block['bbox'][0])/72)))
                except (ValueError,TypeError): pass
                continue
            paragraph=result.add_paragraph()
            for j,line in enumerate(block['lines']):
                if j: paragraph.add_run('\n')
                for span in line['spans']:
                    run=paragraph.add_run(span['text']); run.font.size=Pt(span['size'])
                    run.bold=bool(span['flags']&16); run.italic=bool(span['flags']&2)
                    ink=span.get('color',0); run.font.color.rgb=RGBColor((ink>>16)&255,(ink>>8)&255,ink&255)
    result.save(output)


def _powerpoint(doc,pages,output,options):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    result = Presentation(); result.slide_width = Inches(10)
    result.slide_height = Inches(min(20,max(1,10*doc[pages[0]].rect.height/doc[pages[0]].rect.width)))
    mode = options.get('mode','editable')
    if mode not in ('editable','image'): raise ValueError('Modo inválido: editable ou image.')
    for i in pages:
        page = doc[i]; slide = result.slides.add_slide(result.slide_layouts[6])
        scale=min(result.slide_width/page.rect.width,result.slide_height/page.rect.height)
        left=(result.slide_width-page.rect.width*scale)/2; top=(result.slide_height-page.rect.height*scale)/2
        if mode == 'image':
            slide.shapes.add_picture(io.BytesIO(page.get_pixmap(dpi=150).tobytes('png')),int(left),int(top),width=int(page.rect.width*scale),height=int(page.rect.height*scale))
            continue
        for block in page.get_text('dict')['blocks']:
            x0,y0,x1,y1 = block['bbox']
            transformed=fitz.Rect(block['bbox'])*page.rotation_matrix
            width=max(1,int((x1-x0)*scale)); height=max(1,int((y1-y0)*scale))
            x=int(left+(transformed.x0+transformed.x1)*scale/2-width/2)
            y=int(top+(transformed.y0+transformed.y1)*scale/2-height/2)
            if block['type']==1:
                picture=slide.shapes.add_picture(io.BytesIO(block['image']),x,y,width=width,height=height)
                picture.rotation=page.rotation
            else:
                box = slide.shapes.add_textbox(x,y,width,height)
                box.rotation=page.rotation
                frame = box.text_frame; frame.word_wrap=False
                frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
                for n,line in enumerate(block['lines']):
                    paragraph = frame.paragraphs[0] if n==0 else frame.add_paragraph()
                    for span in line['spans']:
                        run = paragraph.add_run(); run.text=span['text']; run.font.size=Pt(span['size']*scale/12700)
                        run.font.bold=bool(span['flags']&16); run.font.italic=bool(span['flags']&2)
                        ink=span.get('color',0); run.font.color.rgb=RGBColor((ink>>16)&255,(ink>>8)&255,ink&255)
    result.save(output)


def _excel(doc,pages,output):
    from openpyxl import Workbook
    result = Workbook(); result.remove(result.active)
    for i in pages:
        page = doc[i]; sheet = result.create_sheet(f'Página {i+1}'); sheet.freeze_panes='A2'
        tables = page.find_tables().tables
        rows = [row for table in tables for row in table.extract()] if tables else [[line] for line in page.get_text(sort=True).splitlines()]
        for row in rows:
            sheet.append(["'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v for v in row])
    result.save(output)


def _visual_compare(first,second,output):
    from PIL import Image, ImageChops, ImageEnhance
    # The transaction layer owns the staging file; streaming is still atomic to callers.
    with output.open('w',encoding='utf-8') as stream:
        stream.write('<!doctype html><html lang="pt-BR"><meta charset="utf-8"><title>Comparação PDF</title><style>body{font-family:system-ui;margin:24px;background:#f4f5fa;color:#202533}.row{display:flex;gap:16px}figure{margin:0;width:33%}img{width:100%;border:1px solid #ccd}section{margin-bottom:32px}figcaption{padding:8px}</style><h1>Comparação visual de PDFs</h1><p>Renderização a 96 DPI. Diferenças pequenas devem ser revisadas no documento original.</p>')
        for i in range(max(len(first),len(second))):
            images=[]; canvases=[]; diff=None; enhanced=None
            try:
                for document in (first,second):
                    if i<len(document):
                        pix=document[i].get_pixmap(dpi=96,alpha=False)
                        images.append(Image.frombytes('RGB',(pix.width,pix.height),pix.samples))
                        del pix
                    else: images.append(Image.new('RGB',(1,1),'white'))
                size=(max(image.width for image in images),max(image.height for image in images))
                for image in images:
                    if image.size==size: canvases.append(image)
                    else:
                        canvas=Image.new('RGB',size,'white'); canvas.paste(image,(0,0)); canvases.append(canvas)
                diff=ImageChops.difference(*canvases)
                changed=(i>=len(first) or i>=len(second) or first[i].rect!=second[i].rect or diff.getbbox() is not None)
                enhanced=ImageEnhance.Contrast(diff).enhance(3)
                stream.write(f'<section><h2>Página {i+1}: {"alterada" if changed else "igual"}</h2><div class="row">')
                for label,image in zip(('Original','Comparação','Diferenças'),[*canvases,enhanced]):
                    stream.write(f'<figure><img alt="{label}" src="data:image/png;base64,')
                    with io.BytesIO() as data:
                        image.save(data,format='PNG')
                        with data.getbuffer() as buffer: _write_base64(stream,buffer)
                    stream.write(f'"><figcaption>{label}</figcaption></figure>')
                stream.write('</div></section>')
            finally:
                # Release native image allocations before rendering the next page.
                released=set()
                for image in [*images,*canvases,diff,enhanced]:
                    if image is not None and id(image) not in released:
                        released.add(id(image)); image.close()
        stream.write('</html>')


def capabilities():
    office=shutil.which('soffice') or ('/Applications/LibreOffice.app/Contents/MacOS/soffice' if Path('/Applications/LibreOffice.app/Contents/MacOS/soffice').exists() else None)
    return dict(version='0.2.0',operations=list(TOOLS),dependencies=dict(office=bool(office),ocr=bool(shutil.which('ocrmypdf')),tesseract=bool(shutil.which('tesseract')),html=importlib.util.find_spec('playwright') is not None),features=dict(atomic_outputs=True,preserves_pdf_encryption=True,editable_powerpoint=True),limitations=['Layout reconstruction is approximate.','Visual signatures are not cryptographic signatures.','PDF/A requires independent conformance validation.'])


def _cancel_signal(signum, frame):
    raise KeyboardInterrupt


def main():
    signal.signal(signal.SIGTERM,_cancel_signal)
    parser=argparse.ArgumentParser()
    parser.add_argument('operation',nargs='?',choices=[*TOOLS,*ALIASES])
    parser.add_argument('inputs',nargs='*'); parser.add_argument('--output')
    parser.add_argument('--options',default='{}',help='JSON; use PDF_OPTIONS_STDIN=1 para opções sensíveis.')
    parser.add_argument('--capabilities',action='store_true')
    args=parser.parse_args()
    try:
        if args.capabilities:
            print(json.dumps(capabilities(),ensure_ascii=False)); return
        if not args.operation or not args.output: parser.error('operation e --output são obrigatórios.')
        options=json.load(sys.stdin) if os.getenv('PDF_OPTIONS_STDIN')=='1' else json.loads(args.options)
        run(args.operation,args.inputs,args.output,options)
        print(json.dumps(dict(ok=True,output=args.output),ensure_ascii=False))
    except KeyboardInterrupt:
        print(json.dumps(dict(ok=False,error='Operação cancelada.',code='Cancelled'),ensure_ascii=False),file=sys.stderr)
        sys.exit(130)
    except Exception as exc:
        print(json.dumps(dict(ok=False,error=str(exc),code=type(exc).__name__),ensure_ascii=False),file=sys.stderr)
        sys.exit(1)
if __name__=='__main__': main()
