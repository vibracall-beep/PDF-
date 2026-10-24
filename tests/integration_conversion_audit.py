"""Actual Office round trips, rather than inferred conversion fidelity."""
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest

import fitz
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'engine'))
from pdf_engine import run


class ConversionAuditIntegrationTests(unittest.TestCase):
    def setUp(self): self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
    def tearDown(self):
        self.assertFalse(list(self.root.glob('.pdf-output-*')))
        self.assertFalse(list(self.root.glob('.conversion-audit-*')))
        self.temp.cleanup()

    def test_word_roundtrip_inspects_actual_objects_and_equal_rendering(self):
        office=self.root/'converted.docx';document=Document()
        document.add_heading('Conversion audit heading',level=1);document.add_paragraph('Editable paragraph preserved with 123 tokens.')
        table=document.add_table(rows=2,cols=2);table.cell(0,0).text='Product';table.cell(0,1).text='Quantity';table.cell(1,0).text='Apples';table.cell(1,1).text='12';document.save(office)
        original=self.root/'original.pdf';run('office',[office],original)
        original_bytes=original.read_bytes();office_bytes=office.read_bytes();output=self.root/'audit'
        run('conversion-audit',[original,office],output)
        report=json.loads((output/'report.json').read_text())
        self.assertEqual(report['source_sha256'],hashlib.sha256(original_bytes).hexdigest())
        self.assertEqual(report['converted_sha256'],hashlib.sha256(office_bytes).hexdigest())
        self.assertEqual(report['metrics']['source_word_recall'],1)
        self.assertEqual(report['metrics']['pages'][0]['mean_pixel_difference'],0)
        self.assertGreater(report['editable_structure']['counts']['text_runs'],0)
        self.assertEqual(report['editable_structure']['counts']['tables'],1)
        self.assertFalse(report['conformance_claim']);self.assertTrue(report['review_required'])
        self.assertIn('data:image/png;base64,',(output/'comparison.html').read_text())
        self.assertEqual(output.stat().st_mode&0o777,0o700)
        for file in output.iterdir(): self.assertEqual(file.stat().st_mode&0o777,0o600)
        self.assertEqual(original.read_bytes(),original_bytes);self.assertEqual(office.read_bytes(),office_bytes)

    def test_editable_powerpoint_and_excel_are_inspected_after_real_conversion(self):
        fixtures=[]
        slidefile=self.root/'slides.pptx';slides=Presentation();slide=slides.slides.add_slide(slides.slide_layouts[6])
        frame=slide.shapes.add_textbox(100000,100000,5000000,1000000).text_frame;frame.text='Editable presentation text';slides.save(slidefile);fixtures.append((slidefile,'text_runs'))
        workbookfile=self.root/'sheet.xlsx';workbook=Workbook();workbook.active.append(['Product','Quantity']);workbook.active.append(['Apples',12]);workbook.save(workbookfile);workbook.close();fixtures.append((workbookfile,'value_cells'))
        for office,count in fixtures:
            with self.subTest(format=office.suffix):
                original=self.root/(office.stem+'.pdf');run('office',[office],original)
                output=self.root/(office.stem+'-audit');run('conversion-audit',[original,office],output)
                report=json.loads((output/'report.json').read_text())
                self.assertGreater(report['editable_structure']['counts'][count],0)
                self.assertEqual(report['metrics']['source_word_recall'],1)
                self.assertGreater(report['metrics']['converted_word_count'],0)

    def test_changed_office_text_is_detected_even_when_conversion_succeeds(self):
        original=self.root/'original.pdf';office=self.root/'changed.docx'
        with fitz.open() as doc: doc.new_page().insert_text((30,50),'Alpha original retained requirement');doc.save(original)
        document=Document();document.add_paragraph('Entirely different converted words');document.save(office)
        output=self.root/'audit';run('conversion-audit',[original,office],output)
        report=json.loads((output/'report.json').read_text())
        self.assertLess(report['metrics']['source_word_recall'],1)
        self.assertGreater(report['metrics']['pages'][0]['mean_pixel_difference'],0)
        self.assertTrue((output/'roundtrip.pdf').is_file())

    def test_protected_original_needs_password_without_changing_its_bytes(self):
        source=self.root/'protected.pdf';office=self.root/'converted.docx';password='conversion-audit-password'
        with fitz.open() as doc:
            doc.new_page().insert_text((30,50),'Protected original')
            doc.save(source,encryption=fitz.PDF_ENCRYPT_AES_256,user_pw=password,owner_pw='audit-owner')
        document=Document();document.add_paragraph('Protected original');document.save(office)
        before=source.read_bytes()
        with self.assertRaises(ValueError): run('conversion-audit',[source,office],self.root/'denied')
        self.assertFalse((self.root/'denied').exists())
        run('conversion-audit',[source,office],self.root/'audit',{'password':password})
        self.assertEqual(source.read_bytes(),before)
        self.assertNotIn(password,(self.root/'audit'/'report.json').read_text())
