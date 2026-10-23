"""Explicit optional integration suite; requires external engines, never silently skips."""
from pathlib import Path
import csv
import sys
import tempfile
import unittest
import zipfile
import fitz
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'engine'))
from pdf_engine import run
class IntegrationTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.root = Path(self.tmp.name)
    def tearDown(self): self.tmp.cleanup()
    def test_html_render(self):
        source = self.root/'source.html'; source.write_text('<html><body><h1>PDF integration</h1></body></html>')
        out = self.root/'html.pdf'; run('html',[str(source)],out)
        with fitz.open(out) as doc: self.assertIn('PDF integration',doc[0].get_text())
    def test_office(self):
        from docx import Document
        source = self.root/'source.docx'; doc = Document(); doc.add_paragraph('Office integration'); doc.save(source)
        out = self.root/'office.pdf'; run('office',[str(source)],out)
        with fitz.open(out) as doc: self.assertIn('Office integration',doc[0].get_text())
    def test_excel_to_pdf(self):
        from openpyxl import Workbook
        source = self.root / 'source.xlsx'
        book = Workbook()
        book.active.append(['Product', 'Quantity'] + [f'Column {i}' for i in range(3, 31)])
        book.active.append(['Oranges', 42] + ['Test value'] * 27 + ['LastColumnValue'])
        book.save(source)
        book.close()
        out = self.root / 'excel.pdf'
        original = source.read_bytes()
        run('office', [str(source)], out, {'fit_width': True, 'landscape': True})
        self.assertEqual(source.read_bytes(), original)
        with fitz.open(out) as doc:
            self.assertEqual(len(doc), 1)
            self.assertGreater(doc[0].rect.width, doc[0].rect.height)
            text = ''.join(page.get_text() for page in doc)
            self.assertIn('LastColumnValue', text)
            self.assertIn('Oranges', text)
            self.assertIn('42', text)

    def test_powerpoint_to_pdf(self):
        from pptx import Presentation
        source = self.root / 'source.pptx'
        slides = Presentation()
        slide = slides.slides.add_slide(slides.slide_layouts[0])
        slide.shapes.title.text = 'Presentation integration'
        slides.save(source)
        out = self.root / 'powerpoint.pdf'
        run('office', [str(source)], out)
        with fitz.open(out) as doc:
            self.assertEqual(len(doc), 1)
            self.assertIn('Presentation integration', doc[0].get_text())

    def test_odt_to_pdf_preserves_text(self):
        source = self.root / 'source.odt'
        with zipfile.ZipFile(source, 'w') as archive:
            archive.writestr('mimetype', 'application/vnd.oasis.opendocument.text', compress_type=zipfile.ZIP_STORED)
            archive.writestr('META-INF/manifest.xml', '<?xml version="1.0" encoding="UTF-8"?><manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2"><manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/><manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/></manifest:manifest>')
            archive.writestr('content.xml', '<?xml version="1.0" encoding="UTF-8"?><office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2"><office:body><office:text><text:p>ODT integration: ação e coração.</text:p></office:text></office:body></office:document-content>')
        original = source.read_bytes()
        target = self.root / 'odt.pdf'
        run('odt', [source], target)
        with fitz.open(target) as doc:
            text = ''.join(page.get_text() for page in doc)
            self.assertIn('ODT integration', text)
            self.assertIn('coração', text)
        self.assertEqual(source.read_bytes(), original)

    def test_rtf_to_pdf_preserves_unicode_text(self):
        source = self.root / 'source.rtf'
        source.write_text(r'{\rtf1\ansi\deff0{\fonttbl{\f0 Arial;}}\f0\fs24 RTF integration: a\u231?\u227?o.\par}', encoding='ascii')
        original = source.read_bytes()
        target = self.root / 'rtf.pdf'
        run('rtf', [source], target)
        with fitz.open(target) as doc:
            text = ''.join(page.get_text() for page in doc)
            self.assertIn('RTF integration', text)
            self.assertIn('ação', text)
        self.assertEqual(source.read_bytes(), original)

    def test_csv_to_pdf_parses_quoted_cells_without_evaluating_formulas(self):
        source = self.root / 'source.csv'
        tokens = [f'part{index:03d}' for index in range(30)] + ['END_OF_LONG_CELL']
        with source.open('w', encoding='utf-8', newline='') as handle:
            writer = csv.writer(handle)
            writer.writerow(['Product', 'Quantity', 'Note'])
            writer.writerow(['Oranges, fresh', 42, '=1+1'])
            writer.writerow([' '.join(tokens), 84, 'First multiline note\nMiddle multiline note\nLast multiline note'])
        original = source.read_bytes()
        target = self.root / 'csv.pdf'
        run('csv', [source], target, {'fit_width': True, 'landscape': True})
        with fitz.open(target) as doc:
            text = ''.join(page.get_text() for page in doc)
            self.assertIn('Oranges, fresh', text)
            self.assertNotIn('"Oranges, fresh"', text)
            self.assertIn('42', text)
            self.assertIn('=1+1', text)
            for token in tokens:
                self.assertIn(token, text)
            self.assertIn('First multiline note', text)
            self.assertIn('Last multiline note', text)
        self.assertEqual(source.read_bytes(), original)

    def test_ocr_and_pdfa(self):
        source = self.root/'scan.pdf'
        with fitz.open() as doc:
            page = doc.new_page(); page.insert_text((72,100),'HELLO DOCUMENT RECOGNITION',fontsize=24)
            image = page.get_pixmap(dpi=180)
        with fitz.open() as doc:
            page = doc.new_page(); page.insert_image(page.rect,stream=image.tobytes('png')); doc.save(source)
        for operation in ['ocr','pdfa']:
            out = self.root/(operation+'.pdf'); run(operation,[str(source)],out,{'language':'eng'})
            with fitz.open(out) as doc:
                self.assertIn('HELLO',doc[0].get_text())
                if operation == 'pdfa': self.assertIn('pdfaid:part',doc.get_xml_metadata())
if __name__=='__main__': unittest.main()
