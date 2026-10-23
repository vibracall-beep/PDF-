"""Real Tesseract integrations; missing engines fail explicitly, never silently skip."""
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
import fitz
from PIL import Image
from docx import Document
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'engine'))
from pdf_engine import run
from scan_fixture import scanned_form


class StructuredScanIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.source=self.root/'scan.pdf';scanned_form(self.source)
        self.original=self.source.read_bytes();self.options={'language':'eng','source':'auto'}
    def tearDown(self):
        self.assertEqual(self.source.read_bytes(),self.original);self.tmp.cleanup()
    def detect(self,source=None,**options):
        target=self.root/'detected.json';run('form-detect',[source or self.source],target,{**self.options,**options})
        return json.loads(target.read_text())

    def test_raster_grid_becomes_editable_word_and_excel_cells(self):
        expected=[['Product','Amount'],['Apples','12'],['Oranges','34']]
        word=self.root/'scan.docx';excel=self.root/'scan.xlsx'
        run('word',[self.source],word,self.options);run('excel',[self.source],excel,self.options)
        document=Document(word);self.assertEqual(len(document.tables),1)
        self.assertEqual([[cell.text for cell in row.cells] for row in document.tables[0].rows],expected)
        self.assertIn('Order form','\n'.join(paragraph.text for paragraph in document.paragraphs))
        workbook=load_workbook(excel)
        try:self.assertEqual([list(row) for row in workbook.active.iter_rows(values_only=True)],expected)
        finally:workbook.close()

    def test_fields_have_labels_review_scores_and_creatable_coordinates(self):
        result=self.detect();self.assertTrue(result['review_required']);self.assertEqual(result['pages'][0]['source'],'ocr')
        self.assertEqual(len(result['pages'][0]['tables']),1);self.assertEqual(len(result['fields']),3)
        labels={field['label']:field for field in result['fields']}
        self.assertIn('Name:',labels);self.assertIn('Accept terms',labels);self.assertIn('Date:',labels)
        self.assertEqual(labels['Accept terms']['type'],'checkbox')
        for field in result['fields']:
            self.assertIsInstance(field['confidence'],float);self.assertTrue(0<=field['confidence']<=1);self.assertTrue(field['review_required'])
            self.assertGreater(field['rect'][1],200,'Printed O/D glyphs must not become fields.')
        target=self.root/'fillable.pdf';run('form-create',[self.source],target,{'fields':result['fields']})
        with fitz.open(target) as doc:self.assertEqual(len(list(doc[0].widgets())),3)

    def test_rotated_cropped_scan_uses_visible_unrotated_coordinates(self):
        cropped=self.root/'cropped.pdf'
        with fitz.open(self.source) as doc:
            doc[0].set_cropbox(fitz.Rect(20,20,450,500));doc[0].set_rotation(90);doc.save(cropped)
        result=self.detect(cropped)
        name=next(field for field in result['fields'] if field['label']=='Name:')
        for actual,expected in zip(name['rect'],[90,210,320,240]):self.assertAlmostEqual(actual,expected,delta=1)
        target=self.root/'crop-fields.pdf';run('form-create',[cropped],target,{'fields':result['fields']})
        with fitz.open(target) as doc:self.assertEqual(doc[0].rotation,90);self.assertEqual(len(list(doc[0].widgets())),3)

    def test_encrypted_scan_detection_then_creation_keeps_protection(self):
        protected=self.root/'protected.pdf';run('protect',[self.source],protected,{'new_password':'scan-password'})
        result=self.detect(protected,password='scan-password')
        target=self.root/'protected-fields.pdf';run('form-create',[protected],target,{'password':'scan-password','fields':result['fields']})
        with fitz.open(target) as doc:
            self.assertTrue(doc.needs_pass);self.assertTrue(doc.authenticate('scan-password'));self.assertEqual(len(list(doc[0].widgets())),3)

    def test_blank_image_excel_refuses_empty_ocr_and_word_retains_original_image(self):
        blank=self.root/'blank.pdf';data=io.BytesIO();Image.new('RGB',(400,500),'white').save(data,format='PNG')
        with fitz.open() as doc:
            page=doc.new_page(width=400,height=500);page.insert_image(page.rect,stream=data.getvalue());doc.save(blank)
        target=self.root/'blank.xlsx'
        with self.assertRaisesRegex(ValueError,'Nenhum texto'):run('excel',[blank],target,self.options)
        self.assertFalse(target.exists())
        word=self.root/'blank.docx';run('word',[blank],word,self.options)
        self.assertEqual(len(Document(word).inline_shapes),1)

    def test_missing_language_fails_atomically_with_actionable_error(self):
        target=self.root/'missing.json'
        with self.assertRaisesRegex(RuntimeError,'tesseract'):run('form-detect',[self.source],target,{'source':'ocr','language':'zzzz_missing'})
        self.assertFalse(target.exists());self.assertEqual(list(self.root.glob('.pdf-output-*')),[])

    def test_requested_confidence_threshold_marks_review_candidates_without_dropping_text(self):
        result=self.detect(min_confidence=100)
        self.assertTrue(all(field['low_confidence'] for field in result['fields']))
        self.assertGreater(result['pages'][0]['low_confidence_words'],0)
        self.assertIn('Oranges',[word['text'] for word in result['pages'][0]['words']])

    def test_auto_mode_mixed_document_preserves_native_text_and_extracts_scan(self):
        mixed=self.root/'mixed.pdf'
        with fitz.open() as doc,fitz.open(self.source) as scan:
            page=doc.new_page();page.insert_text((40,60),'Native searchable heading');doc.insert_pdf(scan);doc.save(mixed)
        word=self.root/'mixed.docx';run('word',[mixed],word,self.options)
        document=Document(word)
        self.assertIn('Native searchable heading','\n'.join(paragraph.text for paragraph in document.paragraphs))
        self.assertEqual(len(document.tables),1)
        selected=self.root/'selected.xlsx';run('excel',[mixed],selected,{**self.options,'pages':'2'})
        workbook=load_workbook(selected)
        try:self.assertEqual(workbook.sheetnames,['Página 2']);self.assertEqual(workbook.active['A3'].value,'Oranges')
        finally:workbook.close()

if __name__=='__main__':unittest.main()
