"""Real Tesseract evidence for raster merges, borderless rows and deskew."""
from collections import Counter
import hashlib
import json
import math
from pathlib import Path
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile

import fitz
from docx import Document
from openpyxl import load_workbook

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'engine'))
from pdf_engine import run
from scan_analysis import elements
from scan_table_fixture import scanned_table


class ScanTableIntegrationTests(unittest.TestCase):
    def setUp(self): self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
    def tearDown(self): self.temp.cleanup()
    def source(self,**options):
        path=self.root/('scan-'+str(len(list(self.root.glob('scan-*'))))+'.pdf')
        metadata=scanned_table(path,**options);return path,metadata
    def detect(self,source,**options):
        target=self.root/('report-'+str(len(list(self.root.glob('report-*'))))+'.json')
        run('form-detect',[source],target,dict(source='ocr',language='eng',**options))
        return json.loads(target.read_text())['pages'][0]
    def assert_words_retained(self,analysis):
        expected=Counter(word['text'] for word in analysis['words'])
        actual=Counter()
        for entry in elements(analysis):
            if entry['type']==2:
                for row in entry['rows']:
                    for cell in row: actual.update(str(cell or '').split())
            elif entry['type']==0:
                for line in entry['lines']:
                    for span in line['spans']: actual.update(span['text'].split())
        self.assertEqual(actual,expected)

    def test_raster_merged_cells_become_actual_word_and_excel_merges(self):
        source,_=self.source();before=source.read_bytes();analysis=self.detect(source)
        self.assertEqual(len(analysis['tables']),1);table=analysis['tables'][0]
        expected=[['Merged header',None,'Units'],['Group','Apple','10'],[None,'Orange','20']]
        self.assertEqual(table['rows'],expected)
        self.assertEqual({(m['row'],m['column'],m['row_span'],m['column_span']) for m in table['merges']},{(0,0,1,2),(1,0,2,1)})
        self.assert_words_retained(analysis)
        word=self.root/'merged.docx';excel=self.root/'merged.xlsx'
        for operation,target in [('word',word),('excel',excel)]: run(operation,[source],target,dict(language='eng',source='ocr'))
        with zipfile.ZipFile(word) as archive: tree=ET.fromstring(archive.read('word/document.xml'))
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        self.assertTrue(tree.findall('.//w:gridSpan',ns));self.assertTrue(tree.findall('.//w:vMerge',ns))
        text=' '.join(' '.join(element.text or '' for element in tree.findall('.//w:t',ns)).split())
        for value in ('Merged header','Group','Apple','Orange','Inventory report','Review every value'): self.assertIn(value,text)
        book=load_workbook(excel)
        try:
            sheet=book.active;header=next(cell for row in sheet for cell in row if cell.value=='Merged header');r=header.row
            self.assertEqual(set(map(str,sheet.merged_cells.ranges)),{f'A{r}:B{r}',f'A{r+1}:A{r+2}'})
            self.assertEqual([[sheet.cell(r+row,column+1).value for column in range(3)] for row in range(3)],expected)
        finally: book.close()
        self.assertEqual(source.read_bytes(),before)

    def test_borderless_is_opt_in_and_preserves_numeric_cells_and_prose(self):
        source,_=self.source(kind='aligned');default=self.detect(source)
        self.assertEqual(default['tables'],[]);self.assert_words_retained(default)
        aligned=self.detect(source,table_detection='aligned');self.assertEqual(len(aligned['tables']),1)
        self.assertEqual(aligned['tables'][0]['rows'],[['Product','Units','Price'],['Apple','12','1.20'],['Pear','5','3.00']])
        self.assertEqual(aligned['tables'][0]['source'],'ocr-aligned');self.assert_words_retained(aligned)
        word=self.root/'aligned.docx';markdown=self.root/'aligned.md'
        for operation,target in [('word',word),('markdown',markdown)]:run(operation,[source],target,dict(language='eng',source='ocr',table_detection='aligned'))
        self.assertEqual(len(Document(word).tables),1)
        text=markdown.read_text();self.assertIn('| Apple | 12 | 1.20 |',text)
        self.assertIn('Inventory report',text);self.assertIn('Review every value',text)

    def test_broken_interior_rule_is_not_silently_merged_or_discarded(self):
        source,_=self.source(broken=True);analysis=self.detect(source)
        self.assertEqual(analysis['tables'],[]);self.assert_words_retained(analysis)
        self.assertTrue(analysis['words'])

    def test_small_skew_has_inverse_original_geometry_and_retains_merged_cells(self):
        source,metadata=self.source(skew=3.5);before=hashlib.sha256(source.read_bytes()).hexdigest()
        analysis=self.detect(source,ocr_deskew=True)
        self.assertTrue(analysis['preprocessing']['applied'])
        self.assertAlmostEqual(analysis['preprocessing']['rotation_degrees'],-3.5,delta=.3)
        self.assertEqual(len(analysis['tables']),1);table=analysis['tables'][0]
        self.assertEqual(table['rows'],[['Merged header',None,'Units'],['Group','Apple','10'],[None,'Orange','20']])
        self.assertEqual(len(table['merges']),2);self.assert_words_retained(analysis)
        # Independent forward geometry from the known raster fixture rotation.
        scale=metadata['scale'];width,height=metadata['original_size'];new_width,new_height=metadata['size'];angle=math.radians(metadata['skew'])
        for actual,point in zip(table['quad'],[(40,80),(440,80),(440,260),(40,260)]):
            x,y=point[0]*scale-width/2,point[1]*scale-height/2
            expected=((math.cos(angle)*x+math.sin(angle)*y+new_width/2)/scale,(-math.sin(angle)*x+math.cos(angle)*y+new_height/2)/scale)
            self.assertAlmostEqual(actual[0],expected[0],delta=1.2);self.assertAlmostEqual(actual[1],expected[1],delta=1.2)
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(),before)

    def test_deskew_review_apply_keeps_encryption_pixels_and_source_binding(self):
        source,_=self.source(skew=-3);protected=self.root/'protected.pdf';password='deskew-reader'
        run('protect',[source],protected,dict(new_password=password));review_path=self.root/'review.json'
        run('ocr-review',[protected],review_path,dict(password=password,language='eng',ocr_deskew=True))
        review=json.loads(review_path.read_text());self.assertTrue(review['settings']['ocr_deskew'])
        self.assertTrue(review['pages'][0]['preprocessing']['applied']);self.assertEqual(len(review['words'][0]['quad']),4)
        target=self.root/'searchable.pdf';run('ocr-apply',[protected],target,dict(password=password,source_sha256=review['source_sha256'],review_path=str(review_path)))
        with fitz.open(protected) as before,fitz.open(target) as after:
            self.assertTrue(after.needs_pass);after.authenticate(password);before.authenticate(password)
            self.assertEqual(after.permissions,before.permissions)
            self.assertEqual(after[0].get_pixmap(dpi=144).samples,before[0].get_pixmap(dpi=144).samples)
            self.assertIn('Orange',after[0].get_text())
        forged=json.loads(json.dumps(review));forged['words'][0]['quad'][0][0]+=15
        unsigned={key:value for key,value in forged.items() if key!='review_sha256'}
        forged['review_sha256']=hashlib.sha256(json.dumps(unsigned,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()
        rejected=self.root/'forged.pdf'
        with self.assertRaisesRegex(ValueError,'geometria'):
            run('ocr-apply',[protected],rejected,dict(password=password,source_sha256=review['source_sha256'],review=forged))
        self.assertFalse(rejected.exists())

    def test_no_numeric_evidence_and_text_only_skew_do_not_invent_tables(self):
        source,_=self.source(kind='prose',skew=3);analysis=self.detect(source,table_detection='aligned',ocr_deskew=True)
        self.assertTrue(analysis['preprocessing']['applied'])
        self.assertEqual(analysis['preprocessing']['status'],'text-line-projection-evidence')
        self.assertAlmostEqual(analysis['preprocessing']['rotation_degrees'],-3,delta=.3)
        self.assertEqual(analysis['tables'],[]);self.assert_words_retained(analysis)

    def test_markdown_default_keeps_raster_and_unknown_options_refuse_atomically(self):
        source,_=self.source();target=self.root/'default.md';run('markdown',[source],target)
        self.assertIn('data:image/png;base64,',target.read_text())
        for invalid in (dict(ocr_deskew='yes'),dict(ocr_deskew=1),dict(table_detection='guess')):
            target=self.root/'invalid.docx'
            with self.assertRaises(ValueError):run('word',[source],target,dict(source='ocr',language='eng',**invalid))
            self.assertFalse(target.exists())
