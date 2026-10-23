"""Functional regression tests for editing, safety, exports and form workflows."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

import fitz
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'engine'))
from pdf_engine import run


class AdvancedEngineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / 'source.pdf'
        with fitz.open() as doc:
            for i in range(1, 5):
                page = doc.new_page(width=400, height=500)
                page.insert_text((40, 80), f'Page {i} unique content', fontsize=16)
                page.insert_text((40, 120), 'confidential', fontsize=12)
            doc.set_metadata({'title': 'Original title', 'author': 'Author'})
            doc.save(self.source)
        self.original = self.source.read_bytes()

    def tearDown(self):
        self.assertEqual(self.source.read_bytes(), self.original)
        self.temp.cleanup()

    def output(self, name):
        return self.root / name

    def test_range_order_and_duplicate_pages(self):
        target = self.output('ordered.pdf')
        run('organize', [self.source], target, {'pages': '4,1-2,2'})
        with fitz.open(target) as doc:
            self.assertEqual(len(doc), 4)
            self.assertEqual([int(p.get_text().split()[1]) for p in doc], [4, 1, 2, 2])

    def test_page_selection_limits_rotation(self):
        target = self.output('rotated.pdf')
        run('rotate', [self.source], target, {'pages': '2-3', 'degrees': 180})
        with fitz.open(target) as doc:
            self.assertEqual([p.rotation for p in doc], [0, 180, 180, 0])

    def test_page_selection_limits_text_insertion(self):
        target = self.output('text.pdf')
        run('text', [self.source], target, {'pages': '2', 'text': 'SINGLE PAGE MARK'})
        with fitz.open(target) as doc:
            self.assertEqual(['SINGLE PAGE MARK' in p.get_text() for p in doc], [False, True, False, False])

    def test_edit_replaces_existing_text(self):
        target = self.output('edited.pdf')
        run('edit', [self.source], target, {'search': 'confidential', 'replacement': 'public', 'pages': '1'})
        with fitz.open(target) as doc:
            self.assertIn('public', doc[0].get_text())
            self.assertNotIn('confidential', doc[0].get_text())
            self.assertIn('confidential', doc[1].get_text())
            self.assertIn('unique content', doc[0].get_text())

    def test_edit_rejects_replacement_that_cannot_fit(self):
        target = self.output('overflow.pdf')
        with self.assertRaises(ValueError):
            run('edit', [self.source], target, {'search': 'confidential', 'replacement': 'replacement words ' * 100})
        self.assertFalse(target.exists(), 'A replacement that silently disappears must not be published.')

    def test_portuguese_text_survives_save_and_extraction(self):
        target = self.output('portugues.pdf')
        phrase = 'Ação, açúcar e coração'
        run('text', [self.source], target, {'text': phrase, 'pages': '1'})
        with fitz.open(target) as doc:
            self.assertIn(phrase, doc[0].get_text())

    def test_insert_image_and_shape(self):
        image = self.output('sample.png')
        Image.new('RGB', (80, 40), (210, 25, 30)).save(image)
        target = self.output('image.pdf')
        run('image', [self.source], target, {'image': str(image), 'rect': [40, 160, 200, 240], 'pages': '1'})
        with fitz.open(target) as doc:
            self.assertGreater(len(doc[0].get_images()), 0)
            self.assertEqual(len(doc[1].get_images()), 0)
        shaped = self.output('shape.pdf')
        run('shape', [self.source], shaped, {'shape': 'rectangle', 'rect': [40, 160, 200, 240], 'color': [0, 0, 1], 'fill': [1, 0, 0], 'pages': '2'})
        with fitz.open(shaped) as doc:
            self.assertTrue(doc[1].get_drawings())
            self.assertFalse(doc[0].get_drawings())

    def test_create_fill_and_inspect_forms(self):
        form = self.output('form.pdf')
        fields = [
            {'page': 1, 'name': 'customer', 'type': 'text', 'rect': [40, 180, 240, 210]},
            {'page': 1, 'name': 'approved', 'type': 'checkbox', 'rect': [40, 230, 60, 250]},
            {'page': 2, 'name': 'region', 'type': 'choice', 'rect': [40, 180, 240, 210], 'choices': ['North', 'South']},
        ]
        run('form-create', [self.source], form, {'fields': fields})
        filled = self.output('filled.pdf')
        run('form-fill', [form], filled, {'values': {'customer': 'Maria Silva', 'approved': True, 'region': 'South'}})
        with fitz.open(filled) as doc:
            widgets = {w.field_name: w.field_value for page in doc for w in (page.widgets() or [])}
            self.assertEqual(widgets['customer'], 'Maria Silva')
            self.assertEqual(widgets['region'], 'South')
            self.assertNotIn(widgets['approved'], ('Off', '', False, None))
        report = self.output('fields.json')
        run('forms', [filled], report)
        rows = json.loads(report.read_text())
        self.assertEqual({row['name'] for row in rows}, {'customer', 'approved', 'region'})

    def test_editable_powerpoint_has_text_objects(self):
        from pptx import Presentation
        target = self.output('editable.pptx')
        run('powerpoint', [self.source], target, {'mode': 'editable'})
        slides = Presentation(target).slides
        self.assertEqual(len(slides), 4)
        text = '\n'.join(s.text for slide in slides for s in slide.shapes if s.has_text_frame)
        self.assertIn('Page 1 unique content', text)
        self.assertIn('Page 4 unique content', text)

    def test_powerpoint_image_mode_keeps_pages(self):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        target = self.output('image.pptx')
        run('powerpoint', [self.source], target, {'mode': 'image'})
        slides = Presentation(target).slides
        self.assertEqual(len(slides), 4)
        self.assertTrue(all(any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes) for slide in slides))

    def test_visual_compare_detects_graphical_only_change(self):
        changed = self.output('changed.pdf')
        with fitz.open(self.source) as doc:
            doc[0].draw_rect(fitz.Rect(250, 250, 300, 300), color=(1, 0, 0), fill=(1, 0, 0))
            doc.save(changed)
        report = self.output('comparison.html')
        run('compare-visual', [self.source, changed], report)
        html = report.read_text(encoding='utf-8')
        self.assertIn('data:image/png;base64,', html)
        self.assertGreater(len(html), 1000)
        self.assertIn('Página 1: alterada', html)
        self.assertIn('Página 2: igual', html)
        # The input documents have equal text; this report must still render their graphical difference.
        with fitz.open(changed) as a, fitz.open(self.source) as b:
            self.assertEqual(a[0].get_text(), b[0].get_text())
            self.assertNotEqual(a[0].get_pixmap().samples, b[0].get_pixmap().samples)

    def test_capabilities_are_parseable_json(self):
        engine = Path(__file__).resolve().parents[1] / 'engine' / 'pdf_engine.py'
        result = subprocess.run([sys.executable, str(engine), '--capabilities'], capture_output=True, text=True, check=True)
        capabilities = json.loads(result.stdout)
        self.assertIsInstance(capabilities, dict)
        self.assertIn('form-create', json.dumps(capabilities))
        self.assertIn('ocr', json.dumps(capabilities))

    def test_merge_failure_leaves_no_output(self):
        corrupt = self.output('corrupt.pdf')
        corrupt.write_bytes(b'%PDF-1.7\nthis is not a PDF')
        target = self.output('failed.pdf')
        with self.assertRaises(Exception):
            run('merge', [self.source, corrupt], target)
        self.assertFalse(target.exists())

    def test_failed_workflow_leaves_no_output(self):
        target = self.output('failed-workflow.pdf')
        with self.assertRaises(Exception):
            run('workflow', [self.source], target, {'steps': [{'operation': 'rotate'}, {'operation': 'organize', 'options': {'pages': '99'}}]})
        self.assertFalse(target.exists())

    def test_workflow_rejects_non_pdf_intermediate(self):
        target = self.output('invalid-flow.pdf')
        with self.assertRaises(Exception):
            run('workflow', [self.source], target, {'steps': [{'operation': 'markdown'}, {'operation': 'rotate'}]})
        self.assertFalse(target.exists())

    def test_existing_destination_preserved(self):
        target = self.output('existing.pdf')
        sentinel = b'user original bytes'
        target.write_bytes(sentinel)
        with self.assertRaises(Exception):
            run('compress', [self.source], target)
        self.assertEqual(target.read_bytes(), sentinel)

    def test_output_symlink_cannot_overwrite_target(self):
        target = self.output('alias.pdf')
        target.symlink_to(self.source)
        with self.assertRaises(Exception):
            run('rotate', [self.source], target)
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_encrypted_pdf_remains_encrypted_after_rotation(self):
        protected = self.output('protected.pdf')
        password = 'protected-test-123'
        run('protect', [self.source], protected, {'new_password': password})
        rotated = self.output('protected-rotated.pdf')
        run('rotate', [protected], rotated, {'password': password, 'degrees': 90})
        with fitz.open(rotated) as doc:
            self.assertTrue(doc.needs_pass)
            self.assertTrue(doc.authenticate(password))
            self.assertEqual(doc[0].rotation, 90)
            self.assertIn('unique content', doc[0].get_text())

    def test_encrypted_pdf_wrong_password_leaves_no_output(self):
        protected = self.output('protected.pdf')
        run('protect', [self.source], protected, {'new_password': 'protected-test-123'})
        target = self.output('wrong-password.pdf')
        with self.assertRaises(Exception):
            run('rotate', [protected], target, {'password': 'incorrect'})
        self.assertFalse(target.exists())

    def test_workflow_preserves_encryption_across_steps(self):
        protected = self.output('protected.pdf')
        password = 'protected-test-123'
        run('protect', [self.source], protected, {'new_password': password})
        target = self.output('protected-workflow.pdf')
        run('workflow', [protected], target, {'password': password, 'steps': [{'operation': 'rotate'}, {'operation': 'numbers'}]})
        with fitz.open(target) as doc:
            self.assertTrue(doc.needs_pass)
            self.assertTrue(doc.authenticate(password))
            self.assertEqual(doc[0].rotation, 90)

    def test_encrypted_merge_and_split_preserve_protection(self):
        password = 'protected-test-123'
        protected = self.output('protected.pdf')
        run('protect', [self.source], protected, {'new_password': password})
        merged = self.output('merged-protected.pdf')
        run('merge', [protected, self.source], merged, {'password': password})
        with fitz.open(merged) as doc:
            self.assertTrue(doc.needs_pass)
            self.assertTrue(doc.authenticate(password))
            self.assertEqual(len(doc), 8)
        directory = self.output('protected-parts')
        run('split', [protected], directory, {'password': password})
        parts = sorted(directory.glob('*.pdf'))
        self.assertEqual(len(parts), 4)
        for index, part in enumerate(parts, start=1):
            with fitz.open(part) as doc:
                self.assertTrue(doc.needs_pass)
                self.assertTrue(doc.authenticate(password))
                self.assertEqual(len(doc), 1)
                self.assertIn(f'Page {index} unique content', doc[0].get_text())

    def test_redaction_scrubs_metadata_and_attachments(self):
        source = self.output('sensitive.pdf')
        with fitz.open(self.source) as doc:
            doc.set_metadata({'title': 'confidential', 'author': 'confidential'})
            doc.set_xml_metadata('<metadata><secret>confidential</secret></metadata>')
            doc.embfile_add('secret.txt', b'confidential', filename='secret.txt')
            doc.save(source)
        target = self.output('sanitized.pdf')
        run('redact', [source], target, {'text': 'confidential'})
        with fitz.open(target) as doc:
            self.assertNotIn('confidential', ''.join(page.get_text() for page in doc))
            self.assertNotIn('confidential', json.dumps(doc.metadata))
            self.assertEqual(doc.get_xml_metadata(), '')
            self.assertEqual(doc.embfile_count(), 0)
            self.assertIn('unique content', doc[0].get_text())

    def test_cli_invalid_options_returns_structured_failure(self):
        engine = Path(__file__).resolve().parents[1] / 'engine' / 'pdf_engine.py'
        target = self.output('invalid-json.pdf')
        result = subprocess.run([sys.executable, str(engine), 'rotate', str(self.source), '--output', str(target), '--options', '{invalid json'], capture_output=True, text=True)
        self.assertNotEqual(result.returncode, 0)
        payload = json.loads(result.stderr)
        self.assertFalse(payload['ok'])
        self.assertIn('error', payload)
        self.assertNotIn('Traceback', result.stderr)
        self.assertFalse(target.exists())

    def test_excel_formula_text_is_not_executable(self):
        from openpyxl import load_workbook
        source = self.output('formulas.pdf')
        with fitz.open() as doc:
            page = doc.new_page()
            page.insert_text((40, 60), '=HYPERLINK("https://example.invalid")')
            doc.save(source)
        target = self.output('formulas.xlsx')
        run('excel', [source], target)
        book = load_workbook(target)
        try:
            self.assertTrue(book.worksheets[0]['A1'].value.startswith("'="))
            self.assertNotEqual(book.worksheets[0]['A1'].data_type, 'f')
        finally:
            book.close()


if __name__ == '__main__':
    unittest.main()
