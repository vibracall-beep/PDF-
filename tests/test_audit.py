"""Independent adversarial checks of user-visible document safety contracts.

These fixtures inspect reopened outputs, not implementation details. A converter
may explicitly refuse unsupported glyphs, but may never publish corrupt text as
a successful conversion.
"""
import json
import math
import os
from pathlib import Path
import shlex
import subprocess
import sys
import tempfile
import time
import unittest

import fitz

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "engine"))
from pdf_engine import run


class IndependentAuditTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.source = self.root / "original.pdf"
        with fitz.open() as document:
            for index in range(2):
                page = document.new_page(width=500, height=650)
                page.insert_text((40, 100), f"SECRET-781 page {index + 1}", fontsize=24)
                page.insert_text((40, 300), "Public text remains", fontsize=14)
            document.save(self.source)
        self.original_bytes = self.source.read_bytes()

    def tearDown(self):
        self.assertEqual(self.source.read_bytes(), self.original_bytes)
        self.assertEqual(list(self.root.glob(".pdf-output-*")), [], "Staged document data survived the operation.")
        self.temporary.cleanup()

    def test_redaction_removes_sensitive_annotations_and_widgets_inside_area(self):
        source = self.root / "annotated.pdf"
        with fitz.open(self.source) as document:
            page = document[0]
            annotation = page.add_text_annot((45, 80), "SECRET-781 confidential note")
            annotation.set_info(title="SECRET-781", subject="SECRET-781")
            annotation.update()
            for name, value, rect in [
                ("private", "SECRET-781", (40, 120, 260, 150)),
                ("public", "VISIBLE-902", (40, 350, 260, 380)),
            ]:
                widget = fitz.Widget()
                widget.field_name = name
                widget.field_type = fitz.PDF_WIDGET_TYPE_TEXT
                widget.field_value = value
                widget.rect = fitz.Rect(rect)
                page.add_widget(widget)
            page.insert_link({"kind": fitz.LINK_URI, "from": fitz.Rect(40, 155, 200, 180), "uri": "https://example.test/SECRET-781"})
            document.save(source)
        target = self.root / "redacted.pdf"
        run("redact", [source], target, {"pages": "1", "rects": [[30, 30, 310, 210]]})
        with fitz.open(target) as document:
            page = document[0]
            self.assertNotIn("SECRET-781", page.get_text())
            self.assertTrue(all("SECRET-781" not in json.dumps(a.info) for a in page.annots() or []))
            self.assertTrue(all("SECRET-781" not in str(w.field_value) for w in page.widgets() or []))
            self.assertTrue(all("SECRET-781" not in json.dumps(link) for link in page.get_links()))
            self.assertIn("Public text remains", page.get_text())
            self.assertIn("VISIBLE-902", page.get_text())
            self.assertIn("SECRET-781", document[1].get_text(), "Unselected page content must remain intact.")

    def test_unicode_overlay_is_preserved_or_explicitly_refused(self):
        for operation in ("text", "watermark"):
            with self.subTest(operation=operation):
                target = self.root / f"{operation}-unicode.pdf"
                value = "ação Ω 世界"
                try:
                    run(operation, [self.source], target, {"text": value, "font_size": 14, "pages": "1"})
                except ValueError:
                    self.assertFalse(target.exists(), "Unsupported glyphs must fail before publication.")
                else:
                    with fitz.open(target) as document:
                        self.assertIn(value, document[0].get_text(), "Successful output must retain all requested glyphs.")

    def test_unicode_replacement_is_preserved_or_explicitly_refused(self):
        target = self.root / "replacement-unicode.pdf"
        try:
            run("edit", [self.source], target, {"pages": "1", "search": "SECRET-781", "replacement": "你好", "font_size": 14})
        except ValueError:
            self.assertFalse(target.exists())
        else:
            with fitz.open(target) as document:
                self.assertIn("你好", document[0].get_text())
                self.assertNotIn("SECRET-781", document[0].get_text())
                self.assertIn("SECRET-781", document[1].get_text())

    def test_crop_rectangle_is_relative_to_visible_page_even_after_rotation(self):
        for rotation in (0, 90, 180, 270):
            with self.subTest(rotation=rotation):
                source = self.root / f"cropped-{rotation}.pdf"
                with fitz.open() as document:
                    page = document.new_page(width=600, height=800)
                    page.insert_text((150, 180), "KEEP ME", fontsize=20)
                    page.set_cropbox(fitz.Rect(100, 120, 500, 700))
                    page.set_rotation(rotation)
                    document.save(source)
                target = self.root / f"cropped-again-{rotation}.pdf"
                run("crop", [source], target, {"rect": [10, 20, 300, 400]})
                with fitz.open(target) as document:
                    self.assertEqual(document[0].cropbox, fitz.Rect(110, 140, 400, 520))
                    self.assertEqual(document[0].rotation, rotation)
                    self.assertIn("KEEP ME", document[0].get_text())

    def test_repeated_pages_remain_independent_when_later_rotated(self):
        for operation in ("organize", "extract-pages"):
            with self.subTest(operation=operation):
                repeated = self.root / f"{operation}-repeated.pdf"
                target = self.root / f"{operation}-rotated.pdf"
                run(operation, [self.source], repeated, {"pages": "1,1"})
                run("rotate", [repeated], target, {"pages": "1", "degrees": 90})
                with fitz.open(target) as document:
                    self.assertEqual([page.rotation for page in document], [90, 0])

    def test_repeated_pages_remain_independent_when_later_redacted(self):
        for operation in ("organize", "extract-pages"):
            with self.subTest(operation=operation):
                repeated = self.root / f"{operation}-copies.pdf"
                target = self.root / f"{operation}-selective-redaction.pdf"
                run(operation, [self.source], repeated, {"pages": "1,1"})
                run("redact", [repeated], target, {"pages": "1", "text": "SECRET-781"})
                with fitz.open(target) as document:
                    self.assertNotIn("SECRET-781", document[0].get_text())
                    self.assertIn("SECRET-781", document[1].get_text(), "Redacting one copy must not alter its sibling.")

    def test_existing_and_dangling_symlink_destinations_are_never_replaced(self):
        original = self.root / "unrelated.bin"
        original.write_bytes(b"untouched existing data")
        for name, linked in [("existing.pdf", original), ("dangling.pdf", self.root / "absent.pdf")]:
            with self.subTest(name=name):
                destination = self.root / name
                destination.symlink_to(linked)
                with self.assertRaises(ValueError):
                    run("rotate", [self.source], destination)
                self.assertTrue(destination.is_symlink())
                self.assertEqual(destination.readlink(), linked)
        self.assertEqual(original.read_bytes(), b"untouched existing data")
        self.assertFalse((self.root / "absent.pdf").exists())

    def test_input_alias_destination_does_not_overwrite_original(self):
        alias = self.root / "alias.pdf"
        alias.symlink_to(self.source)
        with self.assertRaises(ValueError):
            run("rotate", [alias], self.source)
        self.assertTrue(alias.is_symlink())

    def test_directory_export_refuses_existing_directory_without_changing_contents(self):
        destination = self.root / "pages"
        destination.mkdir()
        (destination / "important.txt").write_text("Keep this file")
        with self.assertRaises(ValueError):
            run("split", [self.source], destination)
        self.assertEqual([p.name for p in destination.iterdir()], ["important.txt"])
        self.assertEqual((destination / "important.txt").read_text(), "Keep this file")

    def test_concurrent_destination_creation_cannot_be_overwritten_at_publication(self):
        """Let a second writer win after conversion starts, for files and dirs."""
        engine = Path(__file__).resolve().parents[1] / "engine" / "pdf_engine.py"
        bin_directory = self.root / "converter bin"
        bin_directory.mkdir()
        worker = bin_directory / "conversion worker.py"
        worker.write_text(
            "import os,sys,time,shutil\n"
            "from pathlib import Path\n"
            "Path(os.environ['AUDIT_STARTED']).write_text('ready')\n"
            "deadline=time.monotonic()+8\n"
            "while not Path(os.environ['AUDIT_RELEASE']).exists():\n"
            " if time.monotonic()>deadline: raise RuntimeError('fixture timeout')\n"
            " time.sleep(.01)\n"
            "destination=Path(sys.argv[sys.argv.index('--outdir')+1])/(Path(sys.argv[-1]).stem+'.pdf')\n"
            "shutil.copyfile(os.environ['AUDIT_PDF'],destination)\n"
        )
        converter = bin_directory / "soffice"
        converter.write_text("#!/bin/sh\nexec " + shlex.quote(sys.executable) + " " + shlex.quote(str(worker)) + ' "$@"\n')
        converter.chmod(0o755)
        source = self.root / "document.docx"
        source.write_bytes(b"stub input: this test concerns publication, not Office fidelity")
        for directory in (False, True):
            with self.subTest(directory=directory):
                destination = self.root / ("concurrent-pages" if directory else "concurrent.pdf")
                started = self.root / f"started-{directory}"
                release = self.root / f"release-{directory}"
                options = {"steps": [{"operation": "office"}, {"operation": "split"}]} if directory else {}
                process = subprocess.Popen(
                    [sys.executable, str(engine), "workflow" if directory else "office", str(source), "--output", str(destination), "--options", json.dumps(options)],
                    stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True,
                    env={**os.environ, "PDF_OPTIONS_STDIN": "0", "PATH": str(bin_directory) + os.pathsep + os.environ.get("PATH", ""),
                         "AUDIT_STARTED": str(started), "AUDIT_RELEASE": str(release), "AUDIT_PDF": str(self.source)},
                )
                try:
                    deadline = time.monotonic() + 6
                    while not started.exists() and process.poll() is None and time.monotonic() < deadline:
                        time.sleep(.01)
                    self.assertTrue(started.exists(), "Converter did not reach the intended race point.")
                    if directory:
                        destination.mkdir()
                        sentinel = destination / "other-user-file.txt"
                    else:
                        sentinel = destination
                    sentinel.write_bytes(b"concurrent writer must win")
                    release.touch()
                    stdout, stderr = process.communicate(timeout=10)
                    self.assertNotEqual(process.returncode, 0, stdout + stderr)
                    self.assertFalse(json.loads(stderr)["ok"])
                    self.assertEqual(sentinel.read_bytes(), b"concurrent writer must win")
                    if directory:
                        self.assertEqual([p.name for p in destination.iterdir()], [sentinel.name])
                finally:
                    release.touch()
                    if process.poll() is None:
                        process.terminate()
                        process.communicate(timeout=10)

    def test_second_invalid_merge_input_never_publishes_partial_document(self):
        corrupt = self.root / "damaged.pdf"
        corrupt.write_bytes(b"not a PDF")
        output = self.root / "partial.pdf"
        with self.assertRaises(Exception):
            run("merge", [self.source, corrupt], output)
        self.assertFalse(output.exists())
        self.assertEqual(corrupt.read_bytes(), b"not a PDF")

    def test_cli_rejects_non_object_json_options_without_publishing(self):
        engine = Path(__file__).resolve().parents[1] / "engine" / "pdf_engine.py"
        for index, invalid in enumerate(([], False, 0, "")):
            with self.subTest(options=invalid):
                target = self.root / f"bad-options-{index}.pdf"
                result = subprocess.run([sys.executable, str(engine), "rotate", str(self.source), "--output", str(target)],
                                        input=json.dumps(invalid), text=True, capture_output=True,
                                        env={**os.environ, "PDF_OPTIONS_STDIN": "1"})
                self.assertNotEqual(result.returncode, 0)
                self.assertFalse(json.loads(result.stderr)["ok"])
                self.assertFalse(target.exists())

    def test_cli_wrong_password_is_not_disclosed_in_json_or_stderr(self):
        password = "correct-pass-9324"
        wrong = "not-the-password-6573"
        protected = self.root / "protected.pdf"
        run("protect", [self.source], protected, {"new_password": password})
        target = self.root / "unauthorized.pdf"
        engine = Path(__file__).resolve().parents[1] / "engine" / "pdf_engine.py"
        result = subprocess.run([sys.executable, str(engine), "unlock", str(protected), "--output", str(target)],
                                input=json.dumps({"password": wrong}), text=True, capture_output=True,
                                env={**os.environ, "PDF_OPTIONS_STDIN": "1"})
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(json.loads(result.stderr)["ok"])
        self.assertNotIn(wrong, result.stdout + result.stderr)
        self.assertNotIn(password, result.stdout + result.stderr)
        self.assertFalse(target.exists())
        with fitz.open(protected) as document:
            self.assertTrue(document.is_encrypted)
            self.assertTrue(document.authenticate(password))

    def test_visual_compare_reports_an_added_blank_page_as_changed(self):
        first = self.root / "one-blank-page.pdf"
        second = self.root / "two-blank-pages.pdf"
        for path, count in ((first, 1), (second, 2)):
            with fitz.open() as document:
                for _ in range(count):
                    document.new_page(width=300, height=400)
                document.save(path)
        report = self.root / "added-page.html"
        run("compare-visual", [first, second], report)
        self.assertIn("Página 1: igual", report.read_text())
        self.assertIn("Página 2: alterada", report.read_text(), "Page presence is part of a visual comparison.")

    def test_visual_compare_reports_changed_page_dimensions(self):
        first = self.root / "small.pdf"
        second = self.root / "large.pdf"
        for path, width, height in ((first, 300, 400), (second, 600, 800)):
            with fitz.open() as document:
                document.new_page(width=width, height=height)
                document.save(path)
        report = self.root / "resized-page.html"
        run("compare-visual", [first, second], report)
        self.assertIn("Página 1: alterada", report.read_text(), "A page size change must not be reported as equal.")

    def test_editable_powerpoint_keeps_rotated_page_text_on_slide(self):
        from pptx import Presentation
        source = self.root / "rotated-source.pdf"
        with fitz.open() as document:
            page = document.new_page(width=500, height=650)
            page.insert_text((40, 600), "Visible near bottom", fontsize=20)
            page.set_rotation(90)
            document.save(source)
        target = self.root / "rotated.pptx"
        run("powerpoint", [source], target)
        presentation = Presentation(target)
        shapes = [shape for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame]
        self.assertEqual([shape.text for shape in shapes], ["Visible near bottom"])
        for shape in shapes:
            # OOXML rotates around the shape center: its unrotated rectangle
            # may cross a slide edge even while every rendered corner fits.
            self.assertEqual(shape.rotation, 90)
            angle = math.radians(shape.rotation)
            center_x = shape.left + shape.width / 2
            center_y = shape.top + shape.height / 2
            for dx in (-shape.width / 2, shape.width / 2):
                for dy in (-shape.height / 2, shape.height / 2):
                    x = center_x + dx * math.cos(angle) - dy * math.sin(angle)
                    y = center_y + dx * math.sin(angle) + dy * math.cos(angle)
                    self.assertGreaterEqual(x, -1)
                    self.assertGreaterEqual(y, -1)
                    self.assertLessEqual(x, presentation.slide_width + 1)
                    self.assertLessEqual(y, presentation.slide_height + 1)

    def make_radio_form(self):
        target = self.root / "radio.pdf"
        run("form-create", [self.source], target, {"fields": [{
            "name": "consent", "type": "radio", "value": "Não",
            "choices": [
                {"value": "Sim", "page": 1, "rect": [40, 400, 60, 420]},
                {"value": "Não", "page": 2, "rect": [40, 400, 60, 420]},
            ],
        }]})
        return target

    def test_radio_group_keeps_single_selection_across_pages_after_round_trip(self):
        source = self.make_radio_form()
        original = source.read_bytes()
        initial = self.root / "radio-initial.json"
        run("forms", [source], initial)
        records = json.loads(initial.read_text())
        self.assertEqual(len(records), 1, "A radio group is one field, not independent checkboxes.")
        self.assertEqual(records[0]["value"], "Não")
        self.assertEqual(records[0]["choices"], ["Sim", "Não"])
        self.assertEqual([button["page"] for button in records[0]["buttons"]], [1, 2])
        target = self.root / "radio-filled.pdf"
        run("form-fill", [source], target, {"values": {"consent": "Sim"}})
        with fitz.open(target) as document:
            states = []
            parents = set()
            for page in document:
                for widget in page.widgets() or []:
                    self.assertEqual(widget.field_type, fitz.PDF_WIDGET_TYPE_RADIOBUTTON)
                    states.append(document.xref_get_key(widget.xref, "AS")[1])
                    parents.add(widget.rb_parent)
            self.assertEqual(states, ["/Sim", "/Off"])
            self.assertEqual(len(parents), 1)
            self.assertEqual(document.xref_get_key(next(iter(parents)), "V"), ("name", "/Sim"))
        self.assertEqual(source.read_bytes(), original)

    def test_radio_group_can_clear_selection_and_rejects_unknown_choice(self):
        source = self.make_radio_form()
        target = self.root / "radio-cleared.pdf"
        run("form-fill", [source], target, {"values": {"consent": ""}})
        with fitz.open(target) as document:
            states = [document.xref_get_key(widget.xref, "AS")[1] for page in document for widget in page.widgets() or []]
            self.assertEqual(states, ["/Off", "/Off"])
        rejected = self.root / "radio-invalid.pdf"
        with self.assertRaises(ValueError):
            run("form-fill", [source], rejected, {"values": {"consent": "Maybe"}})
        self.assertFalse(rejected.exists())

    def test_fill_refuses_read_only_field_without_publishing_partial_update(self):
        source = self.root / "read-only.pdf"
        with fitz.open(self.source) as document:
            page = document[0]
            for index, readonly in enumerate((False, True)):
                widget = fitz.Widget()
                widget.field_name = "locked" if readonly else "editable"
                widget.field_type = fitz.PDF_WIDGET_TYPE_TEXT
                widget.field_value = "Original value"
                widget.field_flags = 1 if readonly else 0
                widget.rect = fitz.Rect(40, 400 + index * 50, 260, 430 + index * 50)
                page.add_widget(widget)
            document.save(source)
        original = source.read_bytes()
        target = self.root / "read-only-filled.pdf"
        with self.assertRaises(ValueError):
            run("form-fill", [source], target, {"values": {"editable": "Changed", "locked": "Unauthorized change"}})
        self.assertFalse(target.exists())
        self.assertEqual(source.read_bytes(), original)

    def test_list_and_dropdown_are_distinct_fields_with_validated_choices(self):
        target = self.root / "choices.pdf"
        fields = [
            {"name": "list", "type": "list", "rect": [40, 400, 220, 470], "choices": ["Ação", "Revisão"], "value": "Revisão"},
            {"name": "dropdown", "type": "choice", "rect": [250, 400, 450, 430], "choices": ["Ação", "Revisão"], "value": "Ação"},
        ]
        run("form-create", [self.source], target, {"fields": fields})
        with fitz.open(target) as document:
            widgets = {widget.field_name: widget for widget in document[0].widgets()}
            self.assertEqual(widgets["list"].field_type, fitz.PDF_WIDGET_TYPE_LISTBOX)
            self.assertEqual(widgets["dropdown"].field_type, fitz.PDF_WIDGET_TYPE_COMBOBOX)
            self.assertEqual(widgets["list"].field_value, "Revisão")
            self.assertEqual(widgets["dropdown"].field_value, "Ação")
        rejected = self.root / "invalid-list-value.pdf"
        with self.assertRaises(ValueError):
            run("form-fill", [target], rejected, {"values": {"list": "Missing option"}})
        self.assertFalse(rejected.exists())


if __name__ == "__main__":
    unittest.main()
