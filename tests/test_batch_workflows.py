"""Real terminal conversions and all-or-nothing workflow batch regressions."""
import contextlib
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

import fitz
from PIL import Image
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'engine'))
import pdf_engine
from pdf_engine import run
from workflow_batch import output_type, result_name, validate_steps


class BatchWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.source = self.root / 'source.pdf'
        self.make_pdf(self.source, 'Alpha selected text')
        self.original = self.source.read_bytes()

    def tearDown(self):
        self.assertEqual(self.source.read_bytes(), self.original)
        self.temporary.cleanup()

    def make_pdf(self, target, text, password=None):
        target.parent.mkdir(parents=True, exist_ok=True)
        with fitz.open() as document:
            for index in range(2):
                page = document.new_page(width=300, height=400)
                page.insert_text((30, 60), text + ' ' + str(index + 1))
            settings = dict(encryption=fitz.PDF_ENCRYPT_AES_256, user_pw=password, owner_pw='owner-' + password) if password else {}
            document.save(target, **settings)

    def steps(self, terminal='markdown'):
        return [{'operation': 'numbers'}, {'operation': terminal, 'options': {'dpi': 72}}]

    def assert_clean(self, destination):
        self.assertFalse(destination.exists())
        self.assertEqual(list(self.root.glob('.pdf-output-*')), [])
        self.assertEqual(list(self.root.glob('.workflow-*')), [])

    def test_every_terminal_export_produces_real_expected_format(self):
        for operation, extension in [('word', 'docx'), ('excel', 'xlsx'), ('powerpoint', 'pptx'), ('markdown', 'md'), ('jpg', ''), ('png', '')]:
            with self.subTest(operation=operation):
                target = self.root / (operation + ('.' + extension if extension else ''))
                run('workflow', [self.source], target, {'steps': self.steps(operation)})
                if extension in ('docx', 'xlsx', 'pptx'):
                    with zipfile.ZipFile(target) as archive:
                        self.assertIn('[Content_Types].xml', archive.namelist())
                    if extension == 'docx':
                        from docx import Document
                        self.assertIn('Alpha', '\n'.join(p.text for p in Document(target).paragraphs))
                    elif extension == 'xlsx':
                        from openpyxl import load_workbook
                        workbook = load_workbook(target)
                        self.assertEqual(len(workbook.sheetnames), 2)
                        workbook.close()
                    else:
                        from pptx import Presentation
                        self.assertEqual(len(Presentation(target).slides), 2)
                elif extension == 'md':
                    self.assertIn('Alpha', target.read_text())
                else:
                    images = sorted(target.glob('*.' + operation))
                    self.assertEqual(len(images), 2)
                    with Image.open(images[0]) as image:
                        self.assertEqual(image.format, 'JPEG' if operation == 'jpg' else 'PNG')
                        self.assertEqual(image.size, (300, 400))

    def test_batch_individual_passwords_duplicate_names_and_secret_free_manifest(self):
        sources = [self.root / 'a' / 'same.pdf', self.root / 'b' / 'same.pdf']
        for path, text, password in zip(sources, ['First private', 'Second private'], ['first-secret', 'second-secret']):
            self.make_pdf(path, text, password)
        output = self.root / 'batch'
        log = io.StringIO()
        options = {'batch': True, 'progress': True, 'input_passwords': ['first-secret', 'second-secret'], 'steps': [{'operation': 'unlock'}, {'operation': 'numbers'}, {'operation': 'word'}]}
        with contextlib.redirect_stdout(log):
            run('workflow', sources, output, options)
        manifest_text = (output / 'batch-manifest.json').read_text()
        self.assertNotIn('secret', manifest_text)
        self.assertNotIn(str(self.root), manifest_text)
        manifest = json.loads(manifest_text)
        self.assertTrue(manifest['atomic'])
        self.assertEqual([item['output'] for item in manifest['results']], ['0001-same.docx', '0002-same.docx'])
        from docx import Document
        self.assertIn('First private', '\n'.join(p.text for p in Document(output / '0001-same.docx').paragraphs))
        self.assertIn('Second private', '\n'.join(p.text for p in Document(output / '0002-same.docx').paragraphs))
        events = [json.loads(line) for line in log.getvalue().splitlines()]
        self.assertEqual(events[0]['completed'], 0)
        self.assertEqual(events[-1]['completed'], 6)
        self.assertTrue(all(event['total'] == 6 for event in events))
        self.assertEqual([event['completed'] for event in events], sorted(event['completed'] for event in events))

    def test_batch_preserves_each_password_across_pdf_stages(self):
        sources = [self.root / 'encrypted1.pdf', self.root / 'encrypted2.pdf']
        credentials = ['password-one', 'password-two']
        for source, password in zip(sources, credentials):
            self.make_pdf(source, 'Protected text', password)
        output = self.root / 'protected-batch'
        run('workflow', sources, output, {'batch': True, 'input_passwords': credentials, 'steps': [{'operation': 'rotate'}, {'operation': 'numbers'}]})
        for index, password in enumerate(credentials):
            with fitz.open(output / f'{index + 1:04d}-encrypted{index + 1}.pdf') as result:
                self.assertTrue(result.needs_pass)
                self.assertTrue(result.authenticate(password))
                self.assertIn('Protected text', result[0].get_text())

    def test_mixed_batch_unlock_then_reprotect_uses_new_credential(self):
        encrypted = self.root / 'encrypted.pdf'
        self.make_pdf(encrypted, 'Private source', 'source-password')
        output = self.root / 'reprotected'
        run('workflow', [self.source, encrypted], output, {
            'batch': True, 'input_passwords': ['', 'source-password'],
            'steps': [{'operation': 'unlock'}, {'operation': 'numbers'}, {'operation': 'protect', 'options': {'new_password': 'new-password'}}],
        })
        for target in output.glob('*.pdf'):
            with fitz.open(target) as document:
                self.assertTrue(document.needs_pass)
                self.assertTrue(document.authenticate('new-password'))
                self.assertIn('1', document[0].get_text())

    def test_batch_nested_image_directories(self):
        output = self.root / 'images'
        run('workflow', [self.source, self.source], output, {'batch': True, 'steps': self.steps('png')})
        self.assertEqual(len(list((output / '0001-source').glob('*.png'))), 2)
        self.assertEqual(len(list((output / '0002-source').glob('*.png'))), 2)

    def test_page_ocr_is_a_standalone_image_to_json_terminal(self):
        tools = set(pdf_engine.TOOLS) | {'ocr-page'}
        self.assertEqual(output_type('ocr-page'), ('json', False))
        self.assertEqual(result_name(Path('scan.png'), 0, 'ocr-page'), '0001-scan.json')
        validate_steps([{'operation': 'ocr-page'}], tools, pdf_engine.PDF_OPERATIONS)
        for steps, batch in [
            ([{'operation': 'rotate'}, {'operation': 'ocr-page'}], False),
            ([{'operation': 'ocr-page'}, {'operation': 'rotate'}], False),
            ([{'operation': 'ocr-page'}], True),
        ]:
            with self.subTest(steps=steps, batch=batch), self.assertRaises(ValueError):
                validate_steps(steps, tools, pdf_engine.PDF_OPERATIONS, batch=batch)

    def test_second_file_failure_publishes_nothing(self):
        encrypted = self.root / 'locked.pdf'
        self.make_pdf(encrypted, 'Private', 'correct-password')
        output = self.root / 'failure'
        with self.assertRaisesRegex(ValueError, '2/2'):
            run('workflow', [self.source, encrypted], output, {'batch': True, 'input_passwords': ['', 'incorrect'], 'steps': self.steps()})
        self.assert_clean(output)

    def test_cancellation_after_first_completed_file_publishes_nothing(self):
        output = self.root / 'cancelled'
        actual_run = run
        calls = 0
        def interrupted_stage(operation, inputs, target, options=None):
            nonlocal calls
            calls += 1
            if calls == 3:
                raise KeyboardInterrupt()
            return actual_run(operation, inputs, target, options)
        with patch.object(pdf_engine, 'run', side_effect=interrupted_stage):
            with self.assertRaises(KeyboardInterrupt):
                actual_run('workflow', [self.source, self.source], output, {'batch': True, 'steps': self.steps()})
        self.assertEqual(calls, 3)
        self.assert_clean(output)

    def test_rejects_wrong_stage_types_and_credentials_before_results(self):
        cases = [
            {'steps': [{'operation': 'word'}, {'operation': 'numbers'}]},
            {'steps': [{'operation': 'numbers'}, {'operation': 'office'}]},
            {'steps': [{'operation': 'workflow'}]},
            {'steps': [{'operation': []}]},
            {'steps': self.steps(), 'batch': 'true'},
            {'steps': self.steps(), 'batch': True, 'input_passwords': ['only-one']},
            {'steps': self.steps(), 'batch': True, 'input_passwords': ['', 1]},
            {'steps': [{'operation': 'merge'}], 'batch': True},
        ]
        for index, options in enumerate(cases):
            output = self.root / f'bad-{index}'
            with self.subTest(index=index), self.assertRaises(ValueError):
                run('workflow', [self.source, self.source], output, options)
            self.assert_clean(output)

    def test_missing_export_dependency_is_reported_before_processing(self):
        output = self.root / 'missing'
        with patch('workflow_batch.importlib.util.find_spec', return_value=None):
            with self.assertRaisesRegex(ValueError, 'docx'):
                run('workflow', [self.source], output, {'steps': self.steps('word')})
        self.assert_clean(output)

    def assert_racing_destination_preserved(self, requested, kind):
        # run() canonicalizes the parent before staging. On macOS /var is an
        # alias of /private/var; compare the same canonical parent without ever
        # following the destination leaf (which may itself become a symlink).
        canonical = requested.parent.resolve() / requested.name
        actual_publish = pdf_engine._publish
        hooks = []
        before = None
        victim = requested.parent / (requested.name + '-link-target')
        victim.write_text('symlink target stays intact')

        def racing_publish(payload, destination):
            nonlocal before
            if Path(destination) == canonical:
                hooks.append(destination)
                if kind == 'empty-directory':
                    requested.mkdir()
                elif kind == 'nonempty-directory':
                    requested.mkdir(); (requested / 'winner').write_text('other process')
                elif kind == 'file':
                    requested.write_text('other process')
                elif kind == 'symlink':
                    requested.symlink_to(victim)
                else:
                    self.fail('Unknown race fixture kind')
                before = requested.lstat()
            actual_publish(payload, destination)

        with patch.object(pdf_engine, '_publish', side_effect=racing_publish):
            try:
                with self.assertRaisesRegex(ValueError, 'Outro processo criou o destino'):
                    run('workflow', [self.source], requested, {'batch': True, 'steps': self.steps()})
            finally:
                self.assertEqual(len(hooks), 1, 'The competing destination must actually be created before publication.')
        after = requested.lstat()
        self.assertEqual((after.st_dev, after.st_ino, after.st_mode), (before.st_dev, before.st_ino, before.st_mode))
        if kind == 'empty-directory':
            self.assertEqual(list(requested.iterdir()), [])
        elif kind == 'nonempty-directory':
            self.assertEqual([p.name for p in requested.iterdir()], ['winner'])
            self.assertEqual((requested / 'winner').read_text(), 'other process')
        elif kind == 'file':
            self.assertEqual(requested.read_text(), 'other process')
        else:
            self.assertTrue(requested.is_symlink())
            self.assertEqual(requested.readlink(), victim)
        self.assertEqual(victim.read_text(), 'symlink target stays intact')
        self.assertEqual(list(requested.parent.glob('.pdf-output-*')), [])
        self.assertEqual(list(requested.parent.glob('.workflow-*')), [])

    def test_existing_and_racing_destinations_never_overwritten(self):
        output = self.root / 'already'
        output.mkdir(); marker = output / 'keep'; marker.write_text('original')
        with self.assertRaises(ValueError):
            run('workflow', [self.source], output, {'batch': True, 'steps': self.steps()})
        self.assertEqual(marker.read_text(), 'original')
        for kind in ('empty-directory', 'nonempty-directory', 'file', 'symlink'):
            with self.subTest(kind=kind):
                self.assert_racing_destination_preserved(self.root / ('race-' + kind), kind)

    def test_racing_destination_with_symlinked_parent_is_really_injected(self):
        # Exercise the macOS /var -> /private/var path shape on every platform,
        # so the test cannot silently omit the race when paths are normalized.
        actual = self.root / 'canonical-parent'; actual.mkdir()
        alias = self.root / 'parent-alias'; alias.symlink_to(actual, target_is_directory=True)
        for kind in ('empty-directory', 'nonempty-directory', 'file', 'symlink'):
            with self.subTest(kind=kind):
                self.assert_racing_destination_preserved(alias / ('race-' + kind), kind)

    def test_safe_names_and_digital_signature_must_be_terminal(self):
        name = result_name('/tmp/.. odd:\nname.pdf', 0, 'word')
        self.assertTrue(name.startswith('0001-'))
        self.assertNotIn(':', name); self.assertNotIn('\n', name); self.assertNotIn('/', name)
        self.assertLess(len(result_name('é' * 200 + '.pdf', 0, 'png').encode()), 200)
        from workflow_batch import validate_steps
        with self.assertRaisesRegex(ValueError, 'assinatura digital'):
            validate_steps([{'operation': 'digital-sign'}, {'operation': 'rotate'}], {'digital-sign', 'rotate'}, {'digital-sign', 'rotate'})

if __name__ == '__main__':
    unittest.main()
