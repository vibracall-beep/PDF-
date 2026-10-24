"""Conversion regressions on actual multi-column and ruled-table PDF content."""
from pathlib import Path
import sys
import tempfile
import unittest

import fitz
from docx import Document
from openpyxl import load_workbook

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'engine'))
from pdf_engine import run
from document_layout import page_elements


class DocumentLayoutTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)

    def tearDown(self): self.temp.cleanup()

    def columns(self):
        source=self.root/'columns.pdf'
        with fitz.open() as doc:
            page=doc.new_page(width=600,height=500)
            page.insert_text((30,45),'A heading spanning both original columns',fontsize=24,fontname='hebo')
            for index in range(4):
                # Interleave insertion order, as many PDF producers do.
                page.insert_text((35,100+index*20),f'Left paragraph {index+1}',fontsize=12)
                page.insert_text((330,100+index*20),f'Right paragraph {index+1}',fontsize=12)
            page.insert_text((30,240),'Full width footer ends the reading sequence.',fontsize=20)
            doc.save(source)
        return source

    def test_word_and_markdown_follow_columns_after_spanning_heading(self):
        source=self.columns();word=self.root/'columns.docx';markdown=self.root/'columns.md'
        run('word',[source],word);run('markdown',[source],markdown)
        document=Document(word)
        text='\n'.join(paragraph.text for paragraph in document.paragraphs)
        self.assertEqual(document.paragraphs[0].style.name,'Heading 1')
        for content in (text,markdown.read_text()):
            positions=[content.index(f'{side} paragraph {index}') for side in ('Left','Right') for index in range(1,5)]
            self.assertEqual(positions,sorted(positions))
            self.assertLess(content.index('heading spanning'),positions[0])
            self.assertGreater(content.index('Full width footer'),positions[-1])

    def table_with_adjacent_text(self):
        source=self.root/'table.pdf'
        with fitz.open() as doc:
            page=doc.new_page(width=500,height=500)
            for x in (100,230,360): page.draw_line((x,50),(x,170))
            for y in (50,90,130,170): page.draw_line((100,y),(360,y))
            for row,values in enumerate([('Product','Units'),('Apples','12'),('Oranges','34')]):
                page.insert_text((10,75+row*40),f'NOTE{row}',fontsize=10)
                for column,value in enumerate(values): page.insert_text((110+column*130,75+row*40),value,fontsize=12)
            page.insert_text((100,210),'Narrative after the table',fontsize=12)
            doc.save(source)
        return source

    def test_adjacent_prose_survives_table_extraction_without_duplicate_cells(self):
        source=self.table_with_adjacent_text();target=self.root/'table.md';word=self.root/'table.docx'
        run('markdown',[source],target);run('word',[source],word)
        content=target.read_text();document=Document(word)
        for label in ('NOTE0','NOTE1','NOTE2','Narrative after the table'): self.assertIn(label,content)
        self.assertEqual(content.count('Apples'),1)
        self.assertEqual(document.tables[0].cell(1,0).text,'Apples')
        self.assertIn('NOTE1','\n'.join(paragraph.text for paragraph in document.paragraphs))

    def test_excel_keeps_prose_table_boundaries_and_cell_provenance(self):
        source=self.table_with_adjacent_text();target=self.root/'table.xlsx';run('excel',[source],target)
        book=load_workbook(target)
        try:
            sheet=book.active;values=[str(cell.value) for row in sheet for cell in row if cell.value is not None]
            self.assertTrue({'NOTE0','NOTE1','NOTE2','Product','Units','Apples','12','Narrative after the table'}.issubset(values))
            headers=[cell for row in sheet for cell in row if cell.value=='Product'];self.assertEqual(len(headers),1)
            self.assertTrue(headers[0].font.bold);self.assertIn('página 1',headers[0].comment.text)
            self.assertTrue(headers[0].alignment.wrap_text);self.assertGreaterEqual(sheet.column_dimensions['A'].width,10)
        finally: book.close()

    def test_markdown_preserves_bold_italic_and_literal_syntax(self):
        source=self.root/'styled.pdf';target=self.root/'styled.md'
        with fitz.open() as doc:
            page=doc.new_page();page.insert_text((30,50),'Bold phrase',fontname='hebo')
            page.insert_text((30,80),'Italic phrase',fontname='heit')
            page.insert_text((30,110),'[literal] *syntax*',fontname='helv');doc.save(source)
        run('markdown',[source],target);content=target.read_text()
        self.assertIn('**Bold phrase**',content);self.assertIn('*Italic phrase*',content)
        self.assertIn(r'\[literal\] \*syntax\*',content)

    def test_reading_order_does_not_omit_overlapping_ambiguous_lines(self):
        source=self.root/'ambiguous.pdf'
        with fitz.open() as doc:
            page=doc.new_page();page.insert_text((30,50),'Alpha overlap');page.insert_text((40,51),'Beta overlap');doc.save(source)
        with fitz.open(source) as doc:
            elements=page_elements(doc[0]);text=''.join(span['text'] for block in elements if block['type']==0 for line in block['lines'] for span in line['spans'])
        self.assertIn('Alpha overlap',text);self.assertIn('Beta overlap',text)

