import json
from pathlib import Path
import sys
import tempfile
import unittest
import fitz
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'engine'))
from pdf_engine import run

class EngineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root/'input.pdf'
        with fitz.open() as doc:
            for text in ['Alpha confidential', 'Beta second page']:
                page = doc.new_page(); page.insert_text((72,100),text)
            doc.save(self.source)
        self.original = self.source.read_bytes()
    def tearDown(self):
        self.assertEqual(self.source.read_bytes(), self.original)
        self.temp.cleanup()
    def target(self,name): return self.root/name
    def test_merge_split_order(self):
        output = self.target('merged.pdf')
        run('merge',[self.source,self.source],output)
        with fitz.open(output) as doc: self.assertEqual(len(doc),4)
        run('split',[output],self.target('split'))
        self.assertEqual(len(list(self.target('split').glob('*.pdf'))),4)
        run('organize',[self.source],self.target('order.pdf'),{'pages':'2,1,2'})
        with fitz.open(self.target('order.pdf')) as doc:
            self.assertEqual(len(doc),3); self.assertIn('Beta',doc[0].get_text())
    def test_encrypt_unlock(self):
        locked = self.target('locked.pdf')
        run('protect',[self.source],locked,{'new_password':'strong-test-123'})
        with fitz.open(locked) as doc:
            self.assertTrue(doc.needs_pass)
        with self.assertRaises(ValueError): run('unlock',[locked],self.target('bad.pdf'),{'password':'wrong'})
        run('unlock',[locked],self.target('plain.pdf'),{'password':'strong-test-123'})
        with fitz.open(self.target('plain.pdf')) as doc:
            self.assertFalse(doc.needs_pass); self.assertIn('Alpha',doc[0].get_text())
    def test_redaction_removes_text(self):
        target = self.target('redacted.pdf')
        run('redact',[self.source],target,{'text':'confidential'})
        with fitz.open(target) as doc:
            self.assertNotIn('confidential',''.join(p.get_text() for p in doc))
            self.assertIn('Alpha',doc[0].get_text())
    def test_rotate_crop_annotations(self):
        run('rotate',[self.source],self.target('rotated.pdf'))
        with fitz.open(self.target('rotated.pdf')) as doc: self.assertEqual(doc[0].rotation,90)
        run('crop',[self.source],self.target('crop.pdf'),{'margin':20})
        with fitz.open(self.target('crop.pdf')) as doc: self.assertEqual(doc[0].cropbox.x0,20)
        run('text',[self.source],self.target('text.pdf'),{'text':'Added'})
        with fitz.open(self.target('text.pdf')) as doc: self.assertIn('Added',doc[0].get_text())
    def test_office_exports(self):
        from docx import Document
        from pptx import Presentation
        from openpyxl import load_workbook
        run('word',[self.source],self.target('out.docx'))
        self.assertIn('Alpha',' '.join(p.text for p in Document(self.target('out.docx')).paragraphs))
        run('powerpoint',[self.source],self.target('out.pptx'))
        self.assertEqual(len(Presentation(self.target('out.pptx')).slides),2)
        run('excel',[self.source],self.target('out.xlsx'))
        book = load_workbook(self.target('out.xlsx'))
        self.assertEqual(len(book.sheetnames),2); self.assertIn('Alpha',book.worksheets[0]['A1'].value)
        book.close()
    def test_images_roundtrip(self):
        run('jpg',[self.source],self.target('jpg'))
        images = sorted(self.target('jpg').glob('*.jpg'))
        self.assertEqual(len(images),2)
        run('images',images,self.target('images.pdf'))
        with fitz.open(self.target('images.pdf')) as doc: self.assertEqual(len(doc),2)
    def test_text_exports_and_compare(self):
        run('markdown',[self.source],self.target('out.md'))
        self.assertIn('Alpha',self.target('out.md').read_text())
        run('compare',[self.source,self.source],self.target('diff.txt'))
        self.assertIn('Nenhuma diferença',self.target('diff.txt').read_text())
        run('forms',[self.source],self.target('forms.json'))
        self.assertEqual(json.loads(self.target('forms.json').read_text()),[])
    def test_workflow(self):
        run('workflow',[self.source],self.target('flow.pdf'),{'steps':[{'operation':'rotate'},{'operation':'numbers'}]})
        with fitz.open(self.target('flow.pdf')) as doc:
            self.assertEqual(doc[0].rotation,90); self.assertIn('1',doc[0].get_text())
    def test_refuse_overwrite_and_invalid_input(self):
        with self.assertRaises(ValueError): run('rotate',[self.source],self.source)
        with self.assertRaises(ValueError): run('organize',[self.source],self.target('invalid.pdf'),{'pages':'99'})
        with self.assertRaises(ValueError): run('redact',[self.source],self.target('empty.pdf'),{})
        with self.assertRaises(ValueError): run('merge',[],self.target('empty.pdf'))
    def test_repair_compress_watermark(self):
        for operation in ['repair','compress','watermark']:
            path = self.target(operation+'.pdf')
            run(operation,[self.source],path,{'text':'WATERMARK'})
            with fitz.open(path) as doc:
                self.assertEqual(len(doc),2); self.assertIn('Alpha',doc[0].get_text())
                if operation == 'watermark': self.assertIn('WATERMARK',doc[0].get_text())
if __name__=='__main__': unittest.main()
