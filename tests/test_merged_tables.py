"""Real ruled/merged and explicit aligned tables, with text-loss checks."""
from collections import Counter
from pathlib import Path
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile

import fitz
from docx import Document
from openpyxl import load_workbook

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'engine'))
from document_layout import page_elements, _table_geometry
from pdf_engine import run


class MergedTableTests(unittest.TestCase):
    def setUp(self): self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
    def tearDown(self): self.temp.cleanup()

    def merged(self):
        source=self.root/'merged.pdf'
        with fitz.open() as doc:
            page=doc.new_page(width=460,height=300)
            page.insert_text((40,25),'Narrative before table')
            for x in (40,280,400): page.draw_line((x,40),(x,160))
            page.draw_line((160,80),(160,160))
            for y in (40,80,160): page.draw_line((40,y),(400,y))
            page.draw_line((160,120),(400,120))
            for x,y,value in [(50,65,'Merged header'),(290,65,'Units'),(50,110,'Group'),(170,105,'Apple'),(290,105,'10'),(170,145,'Orange'),(290,145,'20')]:
                page.insert_text((x,y),value,fontsize=12)
            page.insert_text((40,200),'Narrative after table');doc.save(source)
        return source

    def aligned(self,prose=False):
        source=self.root/('prose.pdf' if prose else 'aligned.pdf')
        rows=[('Product','Units','Price'),('Apple','12','1.20'),('Orange','34','2.00'),('Pear','5','3.50')]
        if prose: rows=[('First story','Other prose','Different text'),('Next phrase','Some content','Extra words'),('Last phrase','Ending prose','Final words')]
        with fitz.open() as doc:
            page=doc.new_page(width=500,height=350);page.insert_text((40,25),'Narrative before aligned data')
            for r,row in enumerate(rows):
                for c,text in enumerate(row): page.insert_text((40+c*130,60+r*30),text,fontsize=12)
            page.insert_text((40,220),'Narrative after aligned data');doc.save(source)
        return source,rows

    def test_geometry_proves_horizontal_and_vertical_spans_without_text_loss(self):
        source=self.merged()
        with fitz.open(source) as doc:
            entries=page_elements(doc[0]);tables=[entry for entry in entries if entry['type']==2]
        self.assertEqual(len(tables),1);table=tables[0]
        self.assertEqual(table['rows'],[['Merged header',None,'Units'],['Group','Apple','10'],[None,'Orange','20']])
        self.assertEqual({(m['row'],m['column'],m['row_span'],m['column_span']) for m in table['merges']},{(0,0,1,2),(1,0,2,1)})
        self.assertEqual(table['column_widths'],[120,120,120])

    def test_word_real_xml_contains_merges_and_every_original_cell_once(self):
        source=self.merged();target=self.root/'merged.docx';run('word',[source],target)
        with zipfile.ZipFile(target) as archive: tree=ET.fromstring(archive.read('word/document.xml'))
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        self.assertTrue(tree.findall('.//w:gridSpan',ns));self.assertTrue(tree.findall('.//w:vMerge',ns))
        text=[node.text for node in tree.findall('.//w:t',ns)]
        for value in ('Merged header','Units','Group','Apple','10','Orange','20','Narrative before table','Narrative after table'):
            self.assertEqual(text.count(value),1,value)
        table=Document(target).tables[0]
        self.assertEqual(table.cell(0,0)._tc,table.cell(0,1)._tc)
        self.assertEqual(table.cell(1,0)._tc,table.cell(2,0)._tc)

    def test_excel_keeps_both_merges_and_original_nonanchor_prose(self):
        source=self.merged();target=self.root/'merged.xlsx';run('excel',[source],target)
        book=load_workbook(target)
        try:
            sheet=book.active;header=next(cell for row in sheet for cell in row if cell.value=='Merged header');r=header.row
            self.assertEqual(set(map(str,sheet.merged_cells.ranges)),{f'A{r}:B{r}',f'A{r+1}:A{r+2}'})
            values=Counter(cell.value for row in sheet for cell in row if cell.value is not None)
            for value in ('Merged header','Units','Group','Apple','10','Orange','20','Narrative before table','Narrative after table'):
                self.assertEqual(values[value],1,value)
            self.assertIsNotNone(header.comment);self.assertTrue(header.alignment.wrap_text)
        finally: book.close()

    def test_ambiguous_grid_cannot_discard_a_covered_nonempty_cell(self):
        class DamagedTable:
            cells=[(0,0,200,20),(0,20,100,40),(100,20,200,40)]
        self.assertIsNone(_table_geometry(DamagedTable(),[['First','must not disappear'],['A','B']]))
        class OverlapTable:
            cells=[(0,0,200,20),(100,0,200,20),(0,20,100,40),(100,20,200,40)]
        self.assertIsNone(_table_geometry(OverlapTable(),[['First',None],['A','B']]))

    def test_borderless_detection_requires_explicit_mode_and_numeric_alignment(self):
        source,rows=self.aligned()
        with fitz.open(source) as doc:
            self.assertFalse(any(entry['type']==2 for entry in page_elements(doc[0])))
            tables=[entry for entry in page_elements(doc[0],dict(table_detection='aligned')) if entry['type']==2]
        self.assertEqual(len(tables),1);self.assertEqual(tables[0]['rows'],[list(row) for row in rows]);self.assertEqual(tables[0]['merges'],[])
        prose,_=self.aligned(prose=True)
        with fitz.open(prose) as doc:
            elements=page_elements(doc[0],dict(table_detection='aligned'))
            self.assertFalse(any(entry['type']==2 for entry in elements))
            text=' '.join(span['text'] for entry in elements if entry['type']==0 for line in entry['lines'] for span in line['spans'])
            self.assertIn('First story',text);self.assertIn('Final words',text)

    def test_borderless_exports_preserve_all_cells_and_surrounding_text(self):
        source,rows=self.aligned();word=self.root/'aligned.docx';excel=self.root/'aligned.xlsx'
        run('word',[source],word,dict(table_detection='aligned'));run('excel',[source],excel,dict(table_detection='aligned'))
        document=Document(word);self.assertEqual([[cell.text for cell in row.cells] for row in document.tables[0].rows],[list(row) for row in rows])
        prose=' '.join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn('Narrative before aligned data',prose);self.assertIn('Narrative after aligned data',prose)
        book=load_workbook(excel)
        try:
            values=[cell.value for row in book.active for cell in row if cell.value is not None]
            for row in rows:
                for value in row: self.assertIn(value,values)
            self.assertIn('Narrative before aligned data',values);self.assertIn('Narrative after aligned data',values)
        finally: book.close()

    def test_unknown_detection_mode_refuses_without_partial_output(self):
        source=self.merged()
        for operation,suffix in [('word','docx'),('excel','xlsx')]:
            target=self.root/('bad.'+suffix)
            with self.assertRaises(ValueError): run(operation,[source],target,dict(table_detection='guess'))
            self.assertFalse(target.exists())

    def test_mixed_ruled_and_borderless_searches_do_not_replace_prior_table_state(self):
        source=self.root/'mixed.pdf'
        first=[['Item','Count'],['Mango','7'],['Plum','8']]
        second=[['Product','Units','Price'],['Pear','12','1.20'],['Peach','5','4.50']]
        with fitz.open() as doc:
            page=doc.new_page(width=500,height=650)
            for x in (40,180,320): page.draw_line((x,40),(x,160))
            for y in (40,80,120,160): page.draw_line((40,y),(320,y))
            for r,row in enumerate(first):
                for c,value in enumerate(row): page.insert_text((50+c*140,65+r*40),value,fontsize=12)
            for r,row in enumerate(second):
                for c,value in enumerate(row): page.insert_text((40+c*130,370+r*30),value,fontsize=12)
            doc.save(source)
        with fitz.open(source) as doc:
            tables=[entry for entry in page_elements(doc[0],dict(table_detection='aligned')) if entry['type']==2]
        self.assertEqual([table['rows'] for table in tables],[first,second])
        target=self.root/'mixed.docx';run('word',[source],target,dict(table_detection='aligned'))
        self.assertEqual([[[cell.text for cell in row.cells] for row in table.rows] for table in Document(target).tables],[first,second])
