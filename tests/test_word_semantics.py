"""Real PDF→DOCX paragraphs, recurring margins and per-page geometry."""
from pathlib import Path
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile

import fitz
from docx import Document

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'engine'))
from pdf_engine import run
from document_layout import _editable_structure


class WordSemanticsTests(unittest.TestCase):
    def setUp(self): self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
    def tearDown(self): self.temp.cleanup()

    def test_repeated_headers_page_labels_and_missing_header_are_separate_sections(self):
        source=self.root/'margins.pdf';target=self.root/'margins.docx'
        with fitz.open() as doc:
            for index in range(3):
                page=doc.new_page(width=500,height=700)
                if index<2: page.insert_text((40,35),'Recurring header',fontsize=10,fontname='hebo')
                page.insert_text((40,140),f'Body of section {index+1}',fontsize=12)
                page.insert_text((230,680),f'Page {index+1} of 3',fontsize=10)
            doc.save(source)
        original=source.read_bytes();run('word',[source],target);result=Document(target)
        self.assertEqual(len(result.sections),3)
        body=' '.join(p.text for p in result.paragraphs)
        self.assertNotIn('Recurring header',body);self.assertNotIn('Page 1 of 3',body)
        for index,section in enumerate(result.sections):
            self.assertAlmostEqual(section.page_width.pt,500);self.assertAlmostEqual(section.page_height.pt,700)
            self.assertIn(f'Body of section {index+1}',body)
            self.assertEqual('\n'.join(p.text for p in section.header.paragraphs),'Recurring header' if index<2 else '')
            self.assertEqual('\n'.join(p.text for p in section.footer.paragraphs),f'Page {index+1} of 3')
            self.assertFalse(section.header.is_linked_to_previous);self.assertFalse(section.footer.is_linked_to_previous)
        self.assertEqual(source.read_bytes(),original)
        structure=_editable_structure(target.read_bytes(),'.docx')
        self.assertEqual(structure['counts']['text_runs'],8)

    def test_single_or_nonrecurrent_margin_text_stays_in_body(self):
        source=self.root/'unique.pdf';target=self.root/'unique.docx'
        with fitz.open() as doc:
            for index in range(2):
                page=doc.new_page(width=500,height=700)
                page.insert_text((40,35),f'Unique title {index+1}',fontsize=12)
                page.insert_text((40,140),f'Body {index+1}',fontsize=12)
                page.insert_text((40,680),f'Unrelated note {index+1}',fontsize=10)
            doc.save(source)
        run('word',[source],target);result=Document(target)
        body=' '.join(p.text for p in result.paragraphs)
        for index in (1,2):
            for label in (f'Unique title {index}',f'Unrelated note {index}'): self.assertIn(label,body)
        self.assertTrue(all(not ''.join(p.text for p in section.header.paragraphs) for section in result.sections))
        selected=self.root/'one.docx';run('word',[source],selected,{'pages':'1'})
        self.assertIn('Unique title 1',' '.join(p.text for p in Document(selected).paragraphs))

    def test_same_margin_words_in_body_or_different_positions_are_not_removed(self):
        source=self.root/'position.pdf';target=self.root/'position.docx'
        with fitz.open() as doc:
            for index in range(3):
                page=doc.new_page(width=500,height=700)
                page.insert_text((40,35),'Running label',fontsize=10)
                page.insert_text((40,140),'Running label',fontsize=10)
                page.insert_text((40+index*70,680),'Different position',fontsize=10)
            doc.save(source)
        run('word',[source],target);result=Document(target)
        body=' '.join(p.text for p in result.paragraphs)
        self.assertEqual(body.count('Running label'),3)
        self.assertEqual(body.count('Different position'),3)
        self.assertTrue(all('Running label' in ''.join(p.text for p in section.header.paragraphs) for section in result.sections))

    def test_consecutive_lines_become_one_editable_paragraph_with_mixed_styles(self):
        source=self.root/'paragraph.pdf';target=self.root/'paragraph.docx'
        with fitz.open() as doc:
            page=doc.new_page(width=500,height=700)
            page.insert_text((40,140),'Regular ',fontsize=12)
            page.insert_text((40+fitz.Font('helv').text_length('Regular ',fontsize=12),140),'bold ending',fontsize=12,fontname='hebo',color=(0.5,0,0))
            page.insert_text((40,156.488),'Next regular line.',fontsize=12)
            page.insert_text((40,220),'Separate paragraph.',fontsize=12)
            doc.save(source)
        run('word',[source],target);result=Document(target)
        paragraphs=[p for p in result.paragraphs if p.text]
        self.assertEqual([p.text for p in paragraphs],['Regular bold ending Next regular line.','Separate paragraph.'])
        bold=next(run for run in paragraphs[0].runs if run.text=='bold ending')
        self.assertTrue(bold.bold);self.assertEqual(str(bold.font.color.rgb),'800000')
        self.assertEqual(bold.font.size.pt,12)

    def test_list_and_column_boundaries_are_not_joined(self):
        source=self.root/'lists.pdf';target=self.root/'lists.docx'
        with fitz.open() as doc:
            page=doc.new_page(width=500,height=700)
            page.insert_text((40,140),'1. First item\n2. Second item',fontsize=12)
            page.insert_text((270,140),'Right first\nRight second',fontsize=12)
            doc.save(source)
        run('word',[source],target);paragraphs=[p.text for p in Document(target).paragraphs if p.text]
        self.assertEqual(paragraphs,['1. First item','2. Second item','Right first Right second'])

    def test_table_widths_and_merges_retain_source_proportions(self):
        source=self.root/'table.pdf';target=self.root/'table.docx'
        with fitz.open() as doc:
            page=doc.new_page(width=500,height=700)
            for x in (40,130,340): page.draw_line((x,100),(x,180))
            for y in (100,140,180): page.draw_line((40,y),(340,y))
            for x,y,text in ((50,125,'ID'),(140,125,'Description'),(50,165,'123'),(140,165,'Example')): page.insert_text((x,y),text)
            doc.save(source)
        run('word',[source],target);table=Document(target).tables[0]
        self.assertFalse(table.autofit)
        self.assertAlmostEqual(table.columns[0].width.pt,90);self.assertAlmostEqual(table.columns[1].width.pt,210)
        self.assertEqual([[cell.text for cell in row.cells] for row in table.rows],[['ID','Description'],['123','Example']])

    def test_office_inspector_counts_margin_text_but_requires_main_body(self):
        target=self.root/'fixture.docx';doc=Document();doc.add_paragraph('Body')
        doc.sections[0].header.paragraphs[0].text='Header';doc.sections[0].footer.paragraphs[0].text='Footer';doc.save(target)
        self.assertEqual(_editable_structure(target.read_bytes(),'.docx')['counts']['text_runs'],3)
        invalid=self.root/'invalid.docx'
        with zipfile.ZipFile(target) as before,zipfile.ZipFile(invalid,'w') as after:
            for name in before.namelist():
                if name!='word/document.xml': after.writestr(name,before.read(name))
        with self.assertRaisesRegex(ValueError,'corpo principal'): _editable_structure(invalid.read_bytes(),'.docx')
